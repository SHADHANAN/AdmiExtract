import io
import os
import openpyxl
import pytest
import pytest_asyncio
from datetime import datetime, timezone
from beanie import PydanticObjectId
from motor.motor_asyncio import AsyncIOMotorClient

from app.core.config import settings
from app.models.user import User, UserRole
from app.models.batch import AdmissionBatch
from app.models.student_submission import StudentSubmission, StudentDocumentMeta
from app.models.student_document import StudentDocument, DocumentVersion
from app.models.wanted_field_config import DocumentFieldConfiguration, WantedFieldItem
from app.models.document_verification import DocumentFieldVerification, VerificationFieldStatus
from app.services.batch_extraction_service import BatchExtractionService
from app.services.excel_template_service import ExcelTemplateService


@pytest_asyncio.fixture(autouse=True)
async def setup_db():
    import app.db.database as db_mod

    db_mod.client = AsyncIOMotorClient(
        settings.MONGODB_URI, serverSelectionTimeoutMS=2000
    )
    db_mod.db = db_mod.client[settings.DATABASE_NAME]
    await db_mod._run_db_setup()


@pytest.mark.asyncio
async def test_excel_template_pipeline_end_to_end(tmp_path):
    """
    ACCEPTANCE TEST (Requirement 24):
    1. Upload template with deliberately unusual headers & column order + unused columns:
       - Column A: Candidate Full Name
       - Column B: REG NO
       - Column C: SEX
       - Column D: Permanent Residence
       - Column E: BANK ACCOUNT
       - Column F: IFSC VALUE
       - Column G: DISTRICT NAME
       - Column H: DOB VALUE
       - Column I: UNUSED_FIELD_1
       - Column J: UNUSED_FIELD_2

    2. Configure field mappings against those exact headers:
       - Student Name -> Candidate Full Name
       - Register Number -> REG NO
       - Gender -> SEX
       - Permanent Address -> Permanent Residence
       - Account Number -> BANK ACCOUNT
       - IFSC Code -> IFSC VALUE
       - District -> DISTRICT NAME
       - Date of Birth -> DOB VALUE

    3. Run extraction for students.
    4. Apply student/staff verification (e.g. approve a correction, confirm another).
    5. Download Verified Excel.
    6. VERIFY:
       - Exact same workbook structure and header order preserved
       - Unused columns (UNUSED_FIELD_1, UNUSED_FIELD_2) remain intact
       - Student Name written into Candidate Full Name (Col A)
       - Register Number written into REG NO (Col B)
       - Gender written into SEX (Col C)
       - Permanent Address written into Permanent Residence (Col D)
       - Account Number written into BANK ACCOUNT (Col E)
       - IFSC Code written into IFSC VALUE (Col F)
       - District written into DISTRICT NAME (Col G)
       - Date of Birth written into DOB VALUE (Col H)
       - 1 student = 1 data row
       - Verified values used correctly
    """
    batch_service = BatchExtractionService()
    template_service = ExcelTemplateService()

    dept_id = f"DEPT_{PydanticObjectId()}"
    batch_id = f"batch_{PydanticObjectId()}"
    batch = AdmissionBatch(
        id=batch_id,
        name="Acceptance_Test_Batch",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch.insert()

    actor = User(
        name="Department Admin",
        username=f"admin_{PydanticObjectId()}",
        email="admin@test.com",
        password="hash",
        role=UserRole.DEPARTMENT_ADMIN,
        department_code=dept_id,
        department_id=dept_id,
        is_active=True,
    )
    await actor.insert()

    # Step 1: Create a custom template workbook with unusual column order & extra columns
    template_headers = [
        "Candidate Full Name",
        "REG NO",
        "SEX",
        "Permanent Residence",
        "BANK ACCOUNT",
        "IFSC VALUE",
        "DISTRICT NAME",
        "DOB VALUE",
        "UNUSED_FIELD_1",
        "UNUSED_FIELD_2",
    ]
    wb_template = openpyxl.Workbook()
    ws_template = wb_template.active
    ws_template.title = "Custom Output Roster"
    ws_template.append(template_headers)

    # Add custom styling to verify preservation
    header_fill = openpyxl.styles.PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    for col_i in range(1, len(template_headers) + 1):
        ws_template.cell(row=1, column=col_i).fill = header_fill

    template_buf = io.BytesIO()
    wb_template.save(template_buf)
    template_bytes = template_buf.getvalue()

    # Step 2: Upload and store template
    preview = await batch_service.parse_and_store_excel_template(
        file_content=template_bytes,
        filename="Student_Details_Columns.xlsx",
        batch_id=batch_id,
        department_id=dept_id,
        uploaded_by=actor.username,
    )
    assert preview.requested_columns == template_headers
    assert preview.column_indices["Candidate Full Name"] == 1
    assert preview.column_indices["REG NO"] == 2
    assert preview.column_indices["SEX"] == 3
    assert preview.column_indices["Permanent Residence"] == 4
    assert preview.column_indices["BANK ACCOUNT"] == 5
    assert preview.column_indices["IFSC VALUE"] == 6
    assert preview.column_indices["DISTRICT NAME"] == 7
    assert preview.column_indices["DOB VALUE"] == 8
    assert preview.column_indices["UNUSED_FIELD_1"] == 9
    assert preview.column_indices["UNUSED_FIELD_2"] == 10

    # Step 3: Configure Wanted Fields for this batch mapped to those exact template headers
    # Aadhaar fields: Student Name, Gender, Date of Birth, Permanent Address, District
    aadhaar_cfg = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="AADHAAR",
        display_name="Aadhaar Card",
        requirement_status="REQUIRED",
        fields=[
            WantedFieldItem(field="Student Name", enabled=True, excel_header="Candidate Full Name"),
            WantedFieldItem(field="Gender", enabled=True, excel_header="SEX"),
            WantedFieldItem(field="Date of Birth", enabled=True, excel_header="DOB VALUE"),
            WantedFieldItem(field="Permanent Address", enabled=True, excel_header="Permanent Residence"),
            WantedFieldItem(field="District", enabled=True, excel_header="DISTRICT NAME"),
        ],
    )
    await aadhaar_cfg.insert()

    # Bank Passbook fields: Account Number, IFSC Code
    passbook_cfg = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="BANK_PASSBOOK",
        display_name="Bank Passbook",
        requirement_status="REQUIRED",
        fields=[
            WantedFieldItem(field="Account Number", enabled=True, excel_header="BANK ACCOUNT"),
            WantedFieldItem(field="IFSC Code", enabled=True, excel_header="IFSC VALUE"),
        ],
    )
    await passbook_cfg.insert()

    # Step 4: Seed 2 Students with documents
    dummy_pdf = str(tmp_path / "doc.pdf")
    with open(dummy_pdf, "wb") as f:
        f.write(b"%PDF-1.4 sample content")

    # Student 1
    sub1 = StudentSubmission(
        batch_id=batch_id,
        register_number="24AM101",
        student_name="Alice Smith",
        mobile_number="9876543210",
        department_id=dept_id,
        documents=[
            StudentDocumentMeta(document_name="Aadhaar Card", document_type="AADHAAR", file_path=dummy_pdf, file_name="aadhaar.pdf"),
            StudentDocumentMeta(document_name="Bank Passbook", document_type="BANK_PASSBOOK", file_path=dummy_pdf, file_name="passbook.pdf"),
        ],
    )
    await sub1.insert()

    doc1_a = StudentDocument(
        student_id="24AM101",
        student_submission_id=str(sub1.id),
        batch_id=batch_id,
        department_id=dept_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        status="ACTIVE",
        active_version=DocumentVersion(
            version_id="v1",
            version_number=1,
            file_path=dummy_pdf,
            file_name="aadhaar.pdf",
            original_filename="aadhaar.pdf",
            file_size_bytes=100,
            mime_type="application/pdf",
            uploaded_by="student",
        ),
    )
    await doc1_a.insert()

    doc1_p = StudentDocument(
        student_id="24AM101",
        student_submission_id=str(sub1.id),
        batch_id=batch_id,
        department_id=dept_id,
        document_type="BANK_PASSBOOK",
        document_name="Bank Passbook",
        status="ACTIVE",
        active_version=DocumentVersion(
            version_id="v1",
            version_number=1,
            file_path=dummy_pdf,
            file_name="passbook.pdf",
            original_filename="passbook.pdf",
            file_size_bytes=100,
            mime_type="application/pdf",
            uploaded_by="student",
        ),
    )
    await doc1_p.insert()

    # Student 2
    sub2 = StudentSubmission(
        batch_id=batch_id,
        register_number="24AM102",
        student_name="Bob Jones",
        mobile_number="9876543211",
        department_id=dept_id,
        documents=[
            StudentDocumentMeta(document_name="Aadhaar Card", document_type="AADHAAR", file_path=dummy_pdf, file_name="aadhaar.pdf"),
            StudentDocumentMeta(document_name="Bank Passbook", document_type="BANK_PASSBOOK", file_path=dummy_pdf, file_name="passbook.pdf"),
        ],
    )
    await sub2.insert()

    doc2_a = StudentDocument(
        student_id="24AM102",
        student_submission_id=str(sub2.id),
        batch_id=batch_id,
        department_id=dept_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        status="ACTIVE",
        active_version=DocumentVersion(
            version_id="v1",
            version_number=1,
            file_path=dummy_pdf,
            file_name="aadhaar.pdf",
            original_filename="aadhaar.pdf",
            file_size_bytes=100,
            mime_type="application/pdf",
            uploaded_by="student",
        ),
    )
    await doc2_a.insert()

    doc2_p = StudentDocument(
        student_id="24AM102",
        student_submission_id=str(sub2.id),
        batch_id=batch_id,
        department_id=dept_id,
        document_type="BANK_PASSBOOK",
        document_name="Bank Passbook",
        status="ACTIVE",
        active_version=DocumentVersion(
            version_id="v1",
            version_number=1,
            file_path=dummy_pdf,
            file_name="passbook.pdf",
            original_filename="passbook.pdf",
            file_size_bytes=100,
            mime_type="application/pdf",
            uploaded_by="student",
        ),
    )
    await doc2_p.insert()

    # Step 5: Execute extraction
    # Register Number is also mapped to REG NO
    column_mapping = {
        "Candidate Full Name": "student_name",
        "REG NO": "register_number",
        "SEX": "gender",
        "Permanent Residence": "permanent_address",
        "BANK ACCOUNT": "account_number",
        "IFSC VALUE": "ifsc_code",
        "DISTRICT NAME": "district",
        "DOB VALUE": "date_of_birth",
    }

    ext_summary = await batch_service.create_and_run_batch_extraction(
        batch_id=batch_id,
        requested_columns=template_headers,
        requested_fields=[
            "student_name",
            "register_number",
            "gender",
            "permanent_address",
            "account_number",
            "ifsc_code",
            "district",
            "date_of_birth",
        ],
        actor=actor,
        column_mapping=column_mapping,
        template_filename="Student_Details_Columns.xlsx",
    )

    req_id = ext_summary.extraction_request_id
    assert req_id is not None

    # Step 6: Simulate Verified Values
    # For Student 1 (24AM101):
    verifs_1 = await DocumentFieldVerification.find(
        DocumentFieldVerification.extraction_request_id == req_id,
        DocumentFieldVerification.student_id == "24AM101",
    ).to_list()

    for v in verifs_1:
        if v.field_name == "student_name":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "Alice Smith Verified"
        elif v.field_name == "gender":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "Female"
        elif v.field_name == "date_of_birth":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "2005-04-12"
        elif v.field_name == "permanent_address":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "123 Main St, Springfield"
        elif v.field_name == "district":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "Coimbatore"
        elif v.field_name == "account_number":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "987654321012"
        elif v.field_name == "ifsc_code":
            v.status = VerificationFieldStatus.STAFF_APPROVED
            v.final_value = "SBIN0001234"
        await v.save()

    # For Student 2 (24AM102):
    verifs_2 = await DocumentFieldVerification.find(
        DocumentFieldVerification.extraction_request_id == req_id,
        DocumentFieldVerification.student_id == "24AM102",
    ).to_list()

    for v in verifs_2:
        if v.field_name == "student_name":
            v.status = VerificationFieldStatus.STUDENT_CONFIRMED
            v.final_value = "Bob Jones"
            v.ai_value = "Bob Jones"
        elif v.field_name == "gender":
            v.status = VerificationFieldStatus.STUDENT_CONFIRMED
            v.final_value = "Male"
            v.ai_value = "Male"
        elif v.field_name == "date_of_birth":
            v.status = VerificationFieldStatus.STUDENT_CONFIRMED
            v.final_value = "2004-11-20"
            v.ai_value = "2004-11-20"
        elif v.field_name == "permanent_address":
            v.status = VerificationFieldStatus.STUDENT_CONFIRMED
            v.final_value = "456 Cross Rd, Salem"
            v.ai_value = "456 Cross Rd, Salem"
        elif v.field_name == "district":
            v.status = VerificationFieldStatus.STUDENT_CONFIRMED
            v.final_value = "Salem"
            v.ai_value = "Salem"
        elif v.field_name == "account_number":
            v.status = VerificationFieldStatus.CORRECTION_REQUESTED  # Pending correction -> NOT exported as final
            v.ai_value = "NOT_AVAILABLE"
            v.student_value = "111122223333"
        elif v.field_name == "ifsc_code":
            v.status = VerificationFieldStatus.STUDENT_CONFIRMED
            v.final_value = "IOBA0005678"
            v.ai_value = "IOBA0005678"
        await v.save()

    # Step 7: Export Verified Excel
    excel_stream, filename = await batch_service.export_verified_excel(req_id, actor)
    assert excel_stream is not None
    assert filename.endswith(".xlsx")

    # Step 8: Open downloaded Excel and strictly inspect columns and values
    wb_out = openpyxl.load_workbook(excel_stream)
    assert "Custom Output Roster" in wb_out.sheetnames
    ws_out = wb_out["Custom Output Roster"]

    # Verify Header Row (Row 1)
    exported_headers = [ws_out.cell(row=1, column=c).value for c in range(1, len(template_headers) + 1)]
    assert exported_headers == template_headers, f"Expected {template_headers}, got {exported_headers}"

    # Verify Exactly 2 Data Rows (Row 2 and Row 3)
    # Row 2 (Student 1: 24AM101)
    row2 = [ws_out.cell(row=2, column=c).value for c in range(1, len(template_headers) + 1)]
    assert row2[0] == "Alice Smith Verified"  # Col A: Candidate Full Name
    assert row2[1] == "24AM101"               # Col B: REG NO
    assert row2[2] == "Female"               # Col C: SEX
    assert row2[3] == "123 Main St, Springfield" # Col D: Permanent Residence
    assert row2[4] == "987654321012"         # Col E: BANK ACCOUNT
    assert row2[5] == "SBIN0001234"          # Col F: IFSC VALUE
    assert row2[6] == "Coimbatore"           # Col G: DISTRICT NAME
    assert row2[7] == "2005-04-12"           # Col H: DOB VALUE
    assert row2[8] is None or row2[8] == ""   # Col I: UNUSED_FIELD_1 (preserved empty)
    assert row2[9] is None or row2[9] == ""   # Col J: UNUSED_FIELD_2 (preserved empty)

    # Row 3 (Student 2: 24AM102)
    row3 = [ws_out.cell(row=3, column=c).value for c in range(1, len(template_headers) + 1)]
    assert row3[0] == "Bob Jones"            # Col A: Candidate Full Name
    assert row3[1] == "24AM102"              # Col B: REG NO
    assert row3[2] == "Male"                 # Col C: SEX
    assert row3[3] == "456 Cross Rd, Salem"  # Col D: Permanent Residence
    assert row3[4] == "" or row3[4] is None  # Col E: BANK ACCOUNT (CORRECTION_REQUESTED / NOT_AVAILABLE -> blank)
    assert row3[5] == "IOBA0005678"          # Col F: IFSC VALUE
    assert row3[6] == "Salem"                # Col G: DISTRICT NAME
    assert row3[7] == "2004-11-20"           # Col H: DOB VALUE
    assert row3[8] is None or row3[8] == ""  # Col I: UNUSED_FIELD_1
    assert row3[9] is None or row3[9] == ""  # Col J: UNUSED_FIELD_2

    # Verify no unexpected rows below Row 3
    for r in range(4, ws_out.max_row + 1):
        for c in range(1, len(template_headers) + 1):
            assert ws_out.cell(row=r, column=c).value is None


@pytest.mark.asyncio
async def test_extraction_history_immutability_and_versioning(tmp_path):
    """
    IMMUTABILITY & VERSIONING TEST (Requirements 14, 15, 16):
    - Extraction #1 uses Template A.
    - Template B is uploaded later for the same batch with completely different headers/order.
    - Extraction #2 uses Template B.
    - Downloading Extraction #1 must STILL output Template A.
    - Downloading Extraction #2 must output Template B.
    """
    batch_service = BatchExtractionService()
    dept_id = f"DEPT_{PydanticObjectId()}"
    batch_id = f"batch_{PydanticObjectId()}"
    batch = AdmissionBatch(
        id=batch_id,
        name="Immutability_Test_Batch",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch.insert()

    actor = User(
        name="Department Admin",
        username=f"admin_{PydanticObjectId()}",
        email="admin@test.com",
        password="hash",
        role=UserRole.DEPARTMENT_ADMIN,
        department_code=dept_id,
        department_id=dept_id,
        is_active=True,
    )
    await actor.insert()

    dummy_pdf = str(tmp_path / "doc.pdf")
    with open(dummy_pdf, "wb") as f:
        f.write(b"%PDF-1.4 sample")

    sub = StudentSubmission(
        batch_id=batch_id,
        register_number="24CS001",
        student_name="Charlie Brown",
        mobile_number="9876543222",
        department_id=dept_id,
        documents=[
            StudentDocumentMeta(document_name="Aadhaar Card", document_type="AADHAAR", file_path=dummy_pdf, file_name="aadhaar.pdf"),
        ],
    )
    await sub.insert()

    doc = StudentDocument(
        student_id="24CS001",
        student_submission_id=str(sub.id),
        batch_id=batch_id,
        department_id=dept_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        status="ACTIVE",
        active_version=DocumentVersion(
            version_id="v1",
            version_number=1,
            file_path=dummy_pdf,
            file_name="aadhaar.pdf",
            original_filename="aadhaar.pdf",
            file_size_bytes=100,
            mime_type="application/pdf",
            uploaded_by="student",
        ),
    )
    await doc.insert()

    # --- TEMPLATE A ---
    headers_a = ["Candidate Name", "Student RegNo", "Gender"]
    wb_a = openpyxl.Workbook()
    ws_a = wb_a.active
    ws_a.title = "Template_A_Sheet"
    ws_a.append(headers_a)
    buf_a = io.BytesIO()
    wb_a.save(buf_a)

    await batch_service.parse_and_store_excel_template(
        file_content=buf_a.getvalue(),
        filename="Template_A.xlsx",
        batch_id=batch_id,
        department_id=dept_id,
        uploaded_by=actor.username,
    )

    aadhaar_cfg = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="AADHAAR",
        display_name="Aadhaar Card",
        requirement_status="REQUIRED",
        fields=[
            WantedFieldItem(field="Student Name", enabled=True, excel_header="Candidate Name"),
            WantedFieldItem(field="Gender", enabled=True, excel_header="Gender"),
        ],
    )
    await aadhaar_cfg.insert()

    ext1 = await batch_service.create_and_run_batch_extraction(
        batch_id=batch_id,
        requested_columns=headers_a,
        requested_fields=["student_name", "register_number", "gender"],
        actor=actor,
        column_mapping={"Candidate Name": "student_name", "Student RegNo": "register_number", "Gender": "gender"},
        template_filename="Template_A.xlsx",
    )
    req1_id = ext1.extraction_request_id

    # --- TEMPLATE B (Uploaded later, replacing Template A) ---
    headers_b = ["REG_NUMBER", "FULL_NAME", "STATE_NAME", "EXTRA_COL"]
    wb_b = openpyxl.Workbook()
    ws_b = wb_b.active
    ws_b.title = "Template_B_Sheet"
    ws_b.append(headers_b)
    buf_b = io.BytesIO()
    wb_b.save(buf_b)

    await batch_service.parse_and_store_excel_template(
        file_content=buf_b.getvalue(),
        filename="Template_B.xlsx",
        batch_id=batch_id,
        department_id=dept_id,
        uploaded_by=actor.username,
    )

    aadhaar_cfg.fields = [
        WantedFieldItem(field="Student Name", enabled=True, excel_header="FULL_NAME"),
        WantedFieldItem(field="State", enabled=True, excel_header="STATE_NAME"),
    ]
    await aadhaar_cfg.save()

    ext2 = await batch_service.create_and_run_batch_extraction(
        batch_id=batch_id,
        requested_columns=headers_b,
        requested_fields=["register_number", "student_name", "state"],
        actor=actor,
        column_mapping={"REG_NUMBER": "register_number", "FULL_NAME": "student_name", "STATE_NAME": "state"},
        template_filename="Template_B.xlsx",
    )
    req2_id = ext2.extraction_request_id

    # Set verified value for Charlie Brown under both extractions
    v1 = await DocumentFieldVerification.find_one(
        DocumentFieldVerification.extraction_request_id == req1_id,
        DocumentFieldVerification.field_name == "student_name",
    )
    if v1:
        v1.status = VerificationFieldStatus.STAFF_APPROVED
        v1.final_value = "Charlie Brown V1"
        await v1.save()

    v2 = await DocumentFieldVerification.find_one(
        DocumentFieldVerification.extraction_request_id == req2_id,
        DocumentFieldVerification.field_name == "student_name",
    )
    if v2:
        v2.status = VerificationFieldStatus.STAFF_APPROVED
        v2.final_value = "Charlie Brown V2"
        await v2.save()

    # --- TEST EXPORT 1 ---
    stream1, fn1 = await batch_service.export_verified_excel(req1_id, actor)
    wb1_out = openpyxl.load_workbook(stream1)
    assert "Template_A_Sheet" in wb1_out.sheetnames
    ws1_out = wb1_out["Template_A_Sheet"]
    exp1_headers = [ws1_out.cell(row=1, column=c).value for c in range(1, len(headers_a) + 1)]
    assert exp1_headers == headers_a
    assert ws1_out.cell(row=2, column=1).value == "Charlie Brown V1"  # Candidate Name
    assert ws1_out.cell(row=2, column=2).value == "24CS001"          # Student RegNo

    # --- TEST EXPORT 2 ---
    stream2, fn2 = await batch_service.export_verified_excel(req2_id, actor)
    wb2_out = openpyxl.load_workbook(stream2)
    assert "Template_B_Sheet" in wb2_out.sheetnames
    ws2_out = wb2_out["Template_B_Sheet"]
    exp2_headers = [ws2_out.cell(row=1, column=c).value for c in range(1, len(headers_b) + 1)]
    assert exp2_headers == headers_b
    assert ws2_out.cell(row=2, column=1).value == "24CS001"          # REG_NUMBER
    assert ws2_out.cell(row=2, column=2).value == "Charlie Brown V2"  # FULL_NAME
    assert ws2_out.cell(row=2, column=4).value is None               # EXTRA_COL preserved
