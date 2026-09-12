from typing import Any, cast
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.models.user import User
from app.db.database import client
from beanie import init_beanie

from motor.motor_asyncio import AsyncIOMotorClient
from app.core.config import settings

@pytest_asyncio.fixture(scope="function", autouse=True)
async def setup_test_db():
    # Monkeypatch AsyncIOMotorClient to bypass Beanie/Motor compatibility issue
    if not hasattr(AsyncIOMotorClient, "append_metadata"):
        AsyncIOMotorClient.append_metadata = lambda *args, **kwargs: None

    test_client = AsyncIOMotorClient(settings.MONGODB_URI)
    test_db = test_client["automate_test_db"]
    
    # Override client and database in app.db.database for the API endpoints
    import app.db.database
    app.db.database.client = test_client
    app.db.database.db = test_db
    app.db.database._is_connected = True

    from app.models.batch import AdmissionBatch
    from app.models.student_submission import StudentSubmission
    from app.models.doc_config_version import DocumentConfigurationVersion
    from app.models.excel_template import ExcelBatchTemplate

    await init_beanie(database=cast(Any, test_db), document_models=[User, AdmissionBatch, StudentSubmission, DocumentConfigurationVersion, ExcelBatchTemplate])
    # Ensure clear collection
    await User.find_all().delete()
    await AdmissionBatch.find_all().delete()
    await StudentSubmission.find_all().delete()
    await DocumentConfigurationVersion.find_all().delete()
    await ExcelBatchTemplate.find_all().delete()
    
    yield
    
    # Clean up test database
    await test_client.drop_database("automate_test_db")
    test_client.close()


@pytest_asyncio.fixture
async def client_async():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        yield ac


@pytest.mark.asyncio
async def test_register_success(client_async: AsyncClient):
    response = await client_async.post(
        "/auth/register",
        json={"username": "testuser", "name": "Test User", "email": "test@example.com", "password": "securepassword123"}
    )
    assert response.status_code == 201
    data = response.json()
    assert data["username"] == "testuser"
    assert data["name"] == "Test User"
    assert data["email"] == "test@example.com"
    assert "id" in data
    assert "password" not in data
    
    # Verify in DB that password is NOT plain text and is hashed
    user = await User.find_one(User.username == "testuser")
    assert user is not None
    assert user.password != "securepassword123"
    assert user.password.startswith("$2b$")  # Bcrypt prefix


@pytest.mark.asyncio
async def test_register_duplicate_email(client_async: AsyncClient):
    # Register first time
    response = await client_async.post(
        "/auth/register",
        json={"username": "testuser1", "name": "Test User", "email": "test@example.com", "password": "securepassword123"}
    )
    assert response.status_code == 201

    # Register second time with same email
    response = await client_async.post(
        "/auth/register",
        json={"username": "testuser2", "name": "Another User", "email": "test@example.com", "password": "anotherpassword"}
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "A user with this email address already exists."


@pytest.mark.asyncio
async def test_login_success(client_async: AsyncClient):
    # Register a user
    await client_async.post(
        "/auth/register",
        json={"username": "testuser", "name": "Test User", "email": "test@example.com", "password": "securepassword123"}
    )

    # Login with username
    response = await client_async.post(
        "/auth/login",
        data={"username": "testuser", "password": "securepassword123"}
    )
    assert response.status_code == 200
    data = response.json()
    assert "access_token" in data
    assert data["token_type"] == "bearer"


@pytest.mark.asyncio
async def test_login_invalid_credentials(client_async: AsyncClient):
    # Register a user
    await client_async.post(
        "/auth/register",
        json={"username": "testuser", "name": "Test User", "email": "test@example.com", "password": "securepassword123"}
    )

    # Login with wrong password
    response = await client_async.post(
        "/auth/login",
        data={"username": "testuser", "password": "wrongpassword"}
    )
    assert response.status_code == 401
    assert response.json()["detail"] == "Incorrect username or password"

    # Login with non-existent username
    response = await client_async.post(
        "/auth/login",
        data={"username": "nonexistent", "password": "securepassword123"}
    )
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_get_current_user(client_async: AsyncClient):
    # Register
    await client_async.post(
        "/auth/register",
        json={"username": "testuser", "name": "Test User", "email": "test@example.com", "password": "securepassword123"}
    )

    # Login to get token
    login_response = await client_async.post(
        "/auth/login",
        data={"username": "testuser", "password": "securepassword123"}
    )
    token = login_response.json()["access_token"]

    # Access /auth/me
    response = await client_async.get(
        "/auth/me",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["username"] == "testuser"
    assert data["email"] == "test@example.com"
    assert data["name"] == "Test User"


@pytest.mark.asyncio
async def test_get_current_user_invalid_token(client_async: AsyncClient):
    response = await client_async.get(
        "/auth/me",
        headers={"Authorization": "Bearer invalidtoken123"}
    )
    assert response.status_code == 401
    assert response.json()["detail"] == "Could not validate credentials"


@pytest.mark.asyncio
async def test_get_students_me(client_async: AsyncClient):
    # Register student user
    await client_async.post(
        "/auth/register",
        json={
            "username": "24AM076",
            "name": "Rahul Sharma",
            "email": "rahul@student.com",
            "password": "studentpassword123",
            "register_number": "24AM076",
            "mobile_number": "9876543210",
            "role": "student"
        }
    )

    # Login as student
    login_response = await client_async.post(
        "/auth/login",
        data={"username": "24AM076", "password": "studentpassword123"}
    )
    assert login_response.status_code == 200
    token = login_response.json()["access_token"]

    # Access /students/me
    response = await client_async.get(
        "/students/me",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["student_name"] == "Rahul Sharma"
    assert data["register_number"] == "24AM076"
    assert data["mobile_number"] == "9876543210"
    assert data["email"] == "rahul@student.com"


@pytest.mark.asyncio
async def test_confirm_submission_security_override(client_async: AsyncClient):
    # Register student user
    await client_async.post(
        "/auth/register",
        json={
            "username": "24AM088",
            "name": "Priya Patel",
            "email": "priya@student.com",
            "password": "priyapassword123",
            "register_number": "24AM088",
            "mobile_number": "9123456789",
            "role": "student"
        }
    )

    # Login as student
    login_response = await client_async.post(
        "/auth/login",
        data={"username": "24AM088", "password": "priyapassword123"}
    )
    token = login_response.json()["access_token"]

    from app.models.batch import AdmissionBatch
    batch = AdmissionBatch(
        id="test_batch_123",
        name="Test Batch",
        academic_year="2025-2026",
        department_id="dept_1",
        created_by="Admin",
        status="active"
    )
    await batch.insert()

    # Try sending tampered values in submission payload
    tampered_payload = {
        "batch_id": "test_batch_123",
        "batch_name": "Test Batch",
        "student_name": "Tampered Name",
        "register_number": "FAKE999",
        "mobile_number": "0000000000",
        "email": "hacker@fake.com",
        "submission_status": "Submitted",
        "documents": [],
        "extracted_data": {}
    }

    # Submit with Authorization header
    response = await client_async.post(
        "/student-submissions/confirm",
        json=tampered_payload,
        headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 200
    data = response.json()

    # Verify backend ignored frontend values and saved DB user values
    assert data["student_name"] == "Priya Patel"
    assert data["register_number"] == "24AM088"
    assert data["mobile_number"] == "9123456789"
    assert data["email"] == "priya@student.com"


@pytest.mark.asyncio
async def test_admin_user_cannot_access_students_me_and_does_not_overwrite_student(client_async: AsyncClient):
    # Register department admin user (AIML Department Admin)
    await client_async.post(
        "/auth/register",
        json={
            "username": "aiml_admin",
            "name": "AIML Department Admin",
            "email": "aiml@test.com",
            "password": "adminpassword123",
            "role": "department_admin"
        }
    )

    # Login as admin
    login_response = await client_async.post(
        "/auth/login",
        data={"username": "aiml_admin", "password": "adminpassword123"}
    )
    admin_token = login_response.json()["access_token"]

    # 1. GET /students/me with Admin token must return 403 Forbidden
    me_response = await client_async.get(
        "/students/me",
        headers={"Authorization": f"Bearer {admin_token}"}
    )
    assert me_response.status_code == 403
    assert "not a student account" in me_response.json()["detail"]

    from app.models.batch import AdmissionBatch
    batch = AdmissionBatch(
        id="test_batch_aiml",
        name="AIML 2025",
        academic_year="2025-2026",
        department_id="aiml_dept",
        created_by="Admin",
        status="active"
    )
    await batch.insert()

    # 2. Student submission with Admin token must NOT overwrite student values with Admin details
    student_payload = {
        "batch_id": "test_batch_aiml",
        "batch_name": "AIML 2025",
        "student_name": "Shadhanan S",
        "register_number": "24AM093",
        "mobile_number": "9488203077",
        "email": "shadhanan@gmail.com",
        "submission_status": "Submitted",
        "documents": [],
        "extracted_data": {}
    }

    sub_response = await client_async.post(
        "/student-submissions/confirm",
        json=student_payload,
        headers={"Authorization": f"Bearer {admin_token}"}
    )
    assert sub_response.status_code == 200
    sub_data = sub_response.json()

    # Confirm it stayed Shadhanan S and was NOT overwritten by AIML Department Admin
    assert sub_data["student_name"] == "Shadhanan S"
    assert sub_data["register_number"] == "24AM093"
    assert sub_data["mobile_number"] == "9488203077"
    assert sub_data["email"] == "shadhanan@gmail.com"


def test_aadhaar_regex_extraction_unit():
    from app.services.ocr_preprocessor import OCRPreprocessor
    preprocessor = OCRPreprocessor()
    
    sample_text = """
    Government of India
    Unique Identification Authority of India
    Name: Rahul Sharma
    DOB: 15/08/2002
    Male
    3849 5012 3491
    """
def test_field_canonicalizer_unit():
    from app.utils.field_canonicalizer import normalize_and_merge_extracted_data, get_canonical_field_name

    assert get_canonical_field_name("Aadhaar Number") == "Aadhaar Card"
    assert get_canonical_field_name("AADHAR CARD") == "Aadhaar Card"
    assert get_canonical_field_name("Caste Certificate") == "Community Certificate"
    assert get_canonical_field_name("Community Code") == "Community Code"
    assert get_canonical_field_name("Community Name") == "Community Name"

    sample = {
        "Aadhaar Card": {"value": None, "confidence": 0},
        "Aadhaar Number": {"value": "9573 0978 4448", "confidence": 100},
        "Aadhaar": {"value": "9573 0978 4448", "confidence": 100},
        "Aadhaar Card Number": {"value": "9573 0978 4448", "confidence": 100},
        "Community Category": {"value": "Backward Class", "confidence": 100},
        "Community Code": {"value": "BC", "confidence": 100},
        "Community Name": {"value": "24 Manai Telugu Chetty", "confidence": 100},
    }

    merged = normalize_and_merge_extracted_data(sample)

    assert "Aadhaar Card" in merged
    assert merged["Aadhaar Card"]["value"] == "9573 0978 4448"
    assert merged["Aadhaar Card"]["confidence"] == 100

    assert "Community Category" in merged
    assert merged["Community Category"]["value"] == "Backward Class"
    assert merged["Community Code"]["value"] == "BC"
    assert merged["Community Name"]["value"] == "24 Manai Telugu Chetty"

    # Verify duplicate keys were eliminated
    assert "Aadhaar Card Number" not in merged
    assert "Community" not in merged


@pytest.mark.asyncio
async def test_append_or_update_student_row_in_excel_unit(tmp_path):
    import openpyxl
    from unittest.mock import AsyncMock
    from app.services.excel_template_service import ExcelTemplateService
    from app.models.excel_template import ExcelBatchTemplate

    # Create dummy template workbook with headers in Row 1
    template_file = tmp_path / "template.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Student Name", "Register Number", "Mobile Number", "Aadhaar Card", "Community"])
    wb.save(template_file)
    wb.close()

    from types import SimpleNamespace
    mock_template = SimpleNamespace(
        batch_id="AIML_2025",
        template_filename="template.xlsx",
        file_path=str(template_file),
        headers=["Student Name", "Register Number", "Mobile Number", "Aadhaar Card", "Community"],
        lookup_column="Register Number",
        field_mappings={},
    )

    mock_repo = AsyncMock()
    mock_repo.get_by_batch_id.return_value = mock_template

    service = ExcelTemplateService(repository=mock_repo)

    student_data = {
        "Student Name": "Praveen",
        "Register Number": "24AM076",
        "Mobile Number": "9488203077",
        "Aadhaar Card": "9573 0978 4448",
        "Community Certificate": "BC",
    }

    # Append first record (Row 2)
    success = await service.append_or_update_student_row_in_excel("AIML_2025", "24AM076", student_data)
    assert success is True

    # Read back saved file
    wb_read = openpyxl.load_workbook(template_file, data_only=True)
    ws_read = wb_read.active
    assert ws_read is not None
    row2 = [cell.value for cell in ws_read[2]]
    assert row2[0] == "Praveen"
    assert row2[1] == "24AM076"
    assert row2[2] == "9488203077"
    assert row2[3] == "9573 0978 4448"
    assert row2[4] == "BC"
    wb_read.close()

    # Update existing record (Row 2)
    student_data_update = {
        "Student Name": "Praveen Kumar",
        "Register Number": "24AM076",
        "Mobile Number": "9488203077",
        "Aadhaar Card": "9573 0978 4448",
        "Community Certificate": "MBC",
    }
    success_update = await service.append_or_update_student_row_in_excel("AIML_2025", "24AM076", student_data_update)
    assert success_update is True

    wb_read2 = openpyxl.load_workbook(template_file, data_only=True)
    ws_read2 = wb_read2.active
    assert ws_read2 is not None
    assert ws_read2.max_row == 2  # No new row created, updated Row 2!
    row2_updated = [cell.value for cell in ws_read2[2]]
    assert row2_updated[0] == "Praveen Kumar"
    assert row2_updated[4] == "MBC"
    wb_read2.close()








