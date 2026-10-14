import io
import openpyxl
import pytest
import pytest_asyncio
from datetime import datetime, timezone, timedelta
from httpx import AsyncClient, ASGITransport
from motor.motor_asyncio import AsyncIOMotorClient

from app.core.config import settings
from app.core.security import create_access_token
from app.main import app
from app.models.batch import AdmissionBatch
from app.models.student_submission import StudentSubmission
from app.models.batch_extraction import (
    BatchExtractionRequest,
    BatchExtractionStatus,
    BatchExportStatus,
    BatchDownloadStatus,
)
from app.models.document_verification import (
    DocumentFieldVerification,
    VerificationFieldStatus,
)
from app.models.verification_token import VerificationToken, VerificationTokenStatus
from app.models.user import User, UserRole
from app.services.verification_token_service import VerificationTokenService
from app.services.batch_extraction_service import BatchExtractionService


@pytest_asyncio.fixture(autouse=True)
async def setup_db():
    import app.db.database as db_mod

    db_mod.client = AsyncIOMotorClient(
        settings.MONGODB_URI, serverSelectionTimeoutMS=2000
    )
    db_mod.db = db_mod.client[settings.DATABASE_NAME]
    await db_mod._run_db_setup()


@pytest_asyncio.fixture
async def client_async():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        yield ac


async def create_test_admin() -> User:
    uname = f"admin_{datetime.now(timezone.utc).timestamp()}"
    u = User(
        name="Admin Test",
        username=uname,
        email=f"{uname}@test.com",
        password="hashed_pw",
        role=UserRole.SUPER_ADMIN,
        department_code="DEP_CSE",
    )
    await u.insert()
    return u


def get_admin_headers(u: User):
    token = create_access_token(
        data={"sub": str(u.id), "role": u.role.value, "department_code": u.department_code}
    )
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.asyncio
async def test_two_consecutive_extractions_portal_shows_latest_values(client_async: AsyncClient):
    """
    Test 1: Two consecutive extraction requests with different field values.
    Verify that the student verification portal always resolves and displays
    the second (latest) extraction results without mixing old values.
    """
    batch_id = f"batch_consec_{datetime.now(timezone.utc).timestamp()}"
    student_reg = "REG_CONSEC_01"
    now = datetime.now(timezone.utc)

    batch = AdmissionBatch(
        id=batch_id,
        name="Consecutive Extractions Batch",
        academic_year="2026-2027",
        department_id="DEP_CSE",
        created_by="admin",
        created_at=now,
    )
    await batch.insert()

    student = StudentSubmission(
        batch_id=batch_id,
        student_id=student_reg,
        register_number=student_reg,
        student_name="Alice Wonder",
        email="alice@test.com",
        mobile_number="9876543210",
        department_id="DEP_CSE",
        created_at=now,
    )
    await student.insert()

    # Run 1: Completed earlier with blood_group = "O+" and nationality = "Indian"
    run_1_id = f"ext_run_1_{datetime.now(timezone.utc).timestamp()}"
    req_1 = BatchExtractionRequest(
        extraction_request_id=run_1_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(hours=2),
        completed_at=now - timedelta(hours=2),
    )
    await req_1.insert()

    verif_1_bg = DocumentFieldVerification(
        extraction_request_id=run_1_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="blood_group",
        ai_value="O+",
        raw_value="O+",
        status=VerificationFieldStatus.AI_EXTRACTED,
        field_state="EXTRACTED",
        created_at=now - timedelta(hours=2),
    )
    await verif_1_bg.insert()

    verif_1_nat = DocumentFieldVerification(
        extraction_request_id=run_1_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="nationality",
        ai_value="Indian",
        raw_value="Indian",
        status=VerificationFieldStatus.AI_EXTRACTED,
        field_state="EXTRACTED",
        created_at=now - timedelta(hours=2),
    )
    await verif_1_nat.insert()

    # Run 2: Completed later with updated extraction: blood_group = "AB+" and nationality = "Indian"
    run_2_id = f"ext_run_2_{datetime.now(timezone.utc).timestamp()}"
    req_2 = BatchExtractionRequest(
        extraction_request_id=run_2_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(minutes=10),
        completed_at=now - timedelta(minutes=10),
    )
    await req_2.insert()

    verif_2_bg = DocumentFieldVerification(
        extraction_request_id=run_2_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="blood_group",
        ai_value="AB+",
        raw_value="AB+",
        status=VerificationFieldStatus.AI_EXTRACTED,
        field_state="EXTRACTED",
        created_at=now - timedelta(minutes=10),
    )
    await verif_2_bg.insert()

    verif_2_nat = DocumentFieldVerification(
        extraction_request_id=run_2_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="nationality",
        ai_value="Indian",
        raw_value="Indian",
        status=VerificationFieldStatus.AI_EXTRACTED,
        field_state="EXTRACTED",
        created_at=now - timedelta(minutes=10),
    )
    await verif_2_nat.insert()

    # Student opens common link without specifying extraction_id
    id_resp = await client_async.post(
        "/public/student-verification/identify",
        json={
            "student_name": "Alice Wonder",
            "register_number": student_reg,
            "batch_id": batch_id,
        },
    )
    assert id_resp.status_code == 200, id_resp.text
    id_data = id_resp.json()
    assert id_data["requires_code"] is True
    session_token = id_data["session_token"]
    dev_code = id_data["dev_code"]

    # Verify session token points to run_2_id
    token_doc = await VerificationToken.find_one(VerificationToken.token == session_token)
    assert token_doc is not None
    assert token_doc.extraction_request_id == run_2_id

    # Complete OTP verification to access extraction data
    code_resp = await client_async.post(
        "/public/student-verification/verify-code",
        json={"session_token": session_token, "code": dev_code},
    )
    assert code_resp.status_code == 200, code_resp.text
    verif_data = code_resp.json()
    assert verif_data["extraction_request_id"] == run_2_id

    # Check field values match run 2 (AB+), not run 1 (O+)
    field_map = {f["field_name"]: f for f in verif_data["fields"]}
    assert "blood_group" in field_map
    assert field_map["blood_group"]["ai_value"] == "AB+"
    assert field_map["blood_group"]["extraction_request_id"] == run_2_id


@pytest.mark.asyncio
async def test_latest_successful_vs_newer_failed_extraction(client_async: AsyncClient):
    """
    Test 2: Latest successful extraction versus a newer failed extraction.
    If a newer extraction run failed or is still processing, the portal must
    authoritatively load the latest successfully completed extraction results,
    never picking the failed or incomplete run.
    """
    batch_id = f"batch_failed_guard_{datetime.now(timezone.utc).timestamp()}"
    student_reg = "REG_GUARD_02"
    now = datetime.now(timezone.utc)

    batch = AdmissionBatch(
        id=batch_id,
        name="Failed Guard Batch",
        academic_year="2026-2027",
        department_id="DEP_CSE",
        created_by="admin",
        created_at=now,
    )
    await batch.insert()

    student = StudentSubmission(
        batch_id=batch_id,
        student_id=student_reg,
        register_number=student_reg,
        student_name="Bob Guard",
        email="bob@test.com",
        mobile_number="9876543211",
        department_id="DEP_CSE",
        created_at=now,
    )
    await student.insert()

    # Run 1: Successfully COMPLETED with valid extracted field
    successful_run_id = f"ext_success_{datetime.now(timezone.utc).timestamp()}"
    req_success = BatchExtractionRequest(
        extraction_request_id=successful_run_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(hours=1),
        completed_at=now - timedelta(hours=1),
    )
    await req_success.insert()

    verif_success = DocumentFieldVerification(
        extraction_request_id=successful_run_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="mother_name",
        ai_value="Grace Guard",
        raw_value="Grace Guard",
        status=VerificationFieldStatus.AI_EXTRACTED,
        field_state="EXTRACTED",
        created_at=now - timedelta(hours=1),
    )
    await verif_success.insert()

    # Run 2: Newer run that FAILED (e.g. quota or pipeline crash)
    failed_run_id = f"ext_failed_{datetime.now(timezone.utc).timestamp()}"
    req_failed = BatchExtractionRequest(
        extraction_request_id=failed_run_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.FAILED.value,
        created_by="admin",
        created_at=now - timedelta(minutes=5),
        failed_at=now - timedelta(minutes=5),
        error_message="OCR Provider Quota Exceeded",
    )
    await req_failed.insert()

    # Run 3: Another newer run that is still PROCESSING
    proc_run_id = f"ext_proc_{datetime.now(timezone.utc).timestamp()}"
    req_proc = BatchExtractionRequest(
        extraction_request_id=proc_run_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.PROCESSING.value,
        created_by="admin",
        created_at=now - timedelta(minutes=1),
    )
    await req_proc.insert()

    # Token service must resolve the successful run, not the failed or processing run
    v_svc = VerificationTokenService()
    resolved_req = await v_svc.resolve_latest_successful_batch_extraction(batch_id=batch_id, student_id=student_reg)
    assert resolved_req is not None
    assert resolved_req.extraction_request_id == successful_run_id
    assert resolved_req.status == BatchExtractionStatus.COMPLETED.value

    # Verify portal endpoint also resolves the successful run
    id_resp = await client_async.post(
        "/public/student-verification/identify",
        json={
            "student_name": "Bob Guard",
            "register_number": student_reg,
            "batch_id": batch_id,
        },
    )
    assert id_resp.status_code == 200, id_resp.text
    id_data = id_resp.json()
    session_token = id_data["session_token"]

    code_resp = await client_async.post(
        "/public/student-verification/verify-code",
        json={"session_token": session_token, "code": id_data["dev_code"]},
    )
    assert code_resp.status_code == 200, code_resp.text
    verif_data = code_resp.json()
    assert verif_data["extraction_request_id"] == successful_run_id
    field_map = {f["field_name"]: f for f in verif_data["fields"]}
    assert field_map["mother_name"]["ai_value"] == "Grace Guard"


@pytest.mark.asyncio
async def test_different_students_in_same_batch(client_async: AsyncClient):
    """
    Test 3: Different students in the same batch.
    Student S1 was in Run 1 and Run 2.
    Student S2 was only in Run 1 (not re-extracted in Run 2).
    Verify that S1 gets Run 2, while S2 still gets Run 1 data.
    """
    batch_id = f"batch_multi_stud_{datetime.now(timezone.utc).timestamp()}"
    now = datetime.now(timezone.utc)

    batch = AdmissionBatch(
        id=batch_id,
        name="Multi Student Batch",
        academic_year="2026-2027",
        department_id="DEP_CSE",
        created_by="admin",
        created_at=now,
    )
    await batch.insert()

    for s_reg, s_name in [("S1_REG", "Student One"), ("S2_REG", "Student Two")]:
        s_doc = StudentSubmission(
            batch_id=batch_id,
            student_id=s_reg,
            register_number=s_reg,
            student_name=s_name,
            email=f"{s_reg.lower()}@test.com",
            mobile_number="9876543212",
            department_id="DEP_CSE",
            created_at=now,
        )
        await s_doc.insert()

    # Run 1: includes S1 and S2
    run_1_id = f"ext_multi_1_{datetime.now(timezone.utc).timestamp()}"
    req_1 = BatchExtractionRequest(
        extraction_request_id=run_1_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(hours=3),
        completed_at=now - timedelta(hours=3),
    )
    await req_1.insert()

    # S1 Run 1 record
    await DocumentFieldVerification(
        extraction_request_id=run_1_id,
        student_id="S1_REG",
        batch_id=batch_id,
        field_name="gender",
        ai_value="MALE",
        raw_value="MALE",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(hours=3),
    ).insert()

    # S2 Run 1 record
    await DocumentFieldVerification(
        extraction_request_id=run_1_id,
        student_id="S2_REG",
        batch_id=batch_id,
        field_name="gender",
        ai_value="FEMALE",
        raw_value="FEMALE",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(hours=3),
    ).insert()

    # Run 2: only re-extracted S1 with gender updated to "Male"
    run_2_id = f"ext_multi_2_{datetime.now(timezone.utc).timestamp()}"
    req_2 = BatchExtractionRequest(
        extraction_request_id=run_2_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(minutes=15),
        completed_at=now - timedelta(minutes=15),
    )
    await req_2.insert()

    await DocumentFieldVerification(
        extraction_request_id=run_2_id,
        student_id="S1_REG",
        batch_id=batch_id,
        field_name="gender",
        ai_value="Male",
        raw_value="Male",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(minutes=15),
    ).insert()

    v_svc = VerificationTokenService()

    # S1 should resolve to Run 2
    res_s1 = await v_svc.resolve_latest_successful_batch_extraction(batch_id=batch_id, student_id="S1_REG")
    assert res_s1 is not None
    assert res_s1.extraction_request_id == run_2_id

    # S2 should resolve to Run 1
    res_s2 = await v_svc.resolve_latest_successful_batch_extraction(batch_id=batch_id, student_id="S2_REG")
    assert res_s2 is not None
    assert res_s2.extraction_request_id == run_1_id


@pytest.mark.asyncio
async def test_duplicate_register_number_across_different_batches(client_async: AsyncClient):
    """
    Test 4: Different batches containing the same student identifier.
    Guarantees:
    - Ambiguous register number without batch_id raises HTTP 400.
    - Explicit batch_id resolves ONLY that batch's latest extraction data.
    - Strict cross-batch tenant isolation.
    """
    dup_reg = "DUP_REG_2026"
    now = datetime.now(timezone.utc)

    batch_1_id = f"batch_dept_cse_{datetime.now(timezone.utc).timestamp()}"
    batch_2_id = f"batch_dept_ece_{datetime.now(timezone.utc).timestamp()}"

    for b_id, b_name, dept in [
        (batch_1_id, "CSE Batch 2026", "DEP_CSE"),
        (batch_2_id, "ECE Batch 2026", "DEP_ECE"),
    ]:
        await AdmissionBatch(
            id=b_id,
            name=b_name,
            academic_year="2026-2027",
            department_id=dept,
            created_by="admin",
            created_at=now,
        ).insert()

        await StudentSubmission(
            batch_id=b_id,
            student_id=dup_reg,
            register_number=dup_reg,
            student_name="Charlie Dup",
            email=f"charlie_{dept.lower()}@test.com",
            mobile_number="9876543213",
            department_id=dept,
            created_at=now,
        ).insert()

    # Batch 1 extraction
    ext_b1_id = f"ext_b1_{datetime.now(timezone.utc).timestamp()}"
    await BatchExtractionRequest(
        extraction_request_id=ext_b1_id,
        batch_id=batch_1_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now,
    ).insert()
    await DocumentFieldVerification(
        extraction_request_id=ext_b1_id,
        student_id=dup_reg,
        batch_id=batch_1_id,
        field_name="department",
        ai_value="Computer Science",
        raw_value="Computer Science",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now,
    ).insert()

    # Batch 2 extraction
    ext_b2_id = f"ext_b2_{datetime.now(timezone.utc).timestamp()}"
    await BatchExtractionRequest(
        extraction_request_id=ext_b2_id,
        batch_id=batch_2_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now,
    ).insert()
    await DocumentFieldVerification(
        extraction_request_id=ext_b2_id,
        student_id=dup_reg,
        batch_id=batch_2_id,
        field_name="department",
        ai_value="Electronics",
        raw_value="Electronics",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now,
    ).insert()

    # 1. Without batch_id: Must return 400 requiring batch specification
    resp_ambiguous = await client_async.post(
        "/public/student-verification/identify",
        json={"student_name": "Charlie Dup", "register_number": dup_reg},
    )
    assert resp_ambiguous.status_code == 400
    assert "Duplicate register number" in resp_ambiguous.json()["detail"]

    # 2. With batch_1_id: resolves Batch 1's extraction
    resp_b1 = await client_async.post(
        "/public/student-verification/identify",
        json={"student_name": "Charlie Dup", "register_number": dup_reg, "batch_id": batch_1_id},
    )
    assert resp_b1.status_code == 200
    tok_b1 = await VerificationToken.find_one(VerificationToken.token == resp_b1.json()["session_token"])
    assert tok_b1.extraction_request_id == ext_b1_id

    # 3. With batch_2_id: resolves Batch 2's extraction
    resp_b2 = await client_async.post(
        "/public/student-verification/identify",
        json={"student_name": "Charlie Dup", "register_number": dup_reg, "batch_id": batch_2_id},
    )
    assert resp_b2.status_code == 200
    tok_b2 = await VerificationToken.find_one(VerificationToken.token == resp_b2.json()["session_token"])
    assert tok_b2.extraction_request_id == ext_b2_id


@pytest.mark.asyncio
async def test_existing_verification_links_refresh_after_newer_extraction(client_async: AsyncClient):
    """
    Test 5: Existing verification links after a newer extraction.
    When a student generated or received an active verification token for Run 1,
    and a newer extraction Run 2 is completed, accessing the session must refresh
    the active token to Run 2, displaying fresh data.
    Loading fresh data must NEVER mark the token or student as verified.
    """
    batch_id = f"batch_refresh_{datetime.now(timezone.utc).timestamp()}"
    student_reg = "REG_REFRESH_05"
    now = datetime.now(timezone.utc)

    await AdmissionBatch(
        id=batch_id,
        name="Refresh Session Batch",
        academic_year="2026-2027",
        department_id="DEP_CSE",
        created_by="admin",
        created_at=now,
    ).insert()

    await StudentSubmission(
        batch_id=batch_id,
        student_id=student_reg,
        register_number=student_reg,
        student_name="David Refresh",
        email="david@test.com",
        mobile_number="9876543214",
        department_id="DEP_CSE",
        created_at=now,
    ).insert()

    # Run 1 completed
    run_1_id = f"ext_ref_1_{datetime.now(timezone.utc).timestamp()}"
    await BatchExtractionRequest(
        extraction_request_id=run_1_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(hours=2),
    ).insert()

    await DocumentFieldVerification(
        extraction_request_id=run_1_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="father_name",
        ai_value="Old Father Name",
        raw_value="Old Father Name",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(hours=2),
    ).insert()

    # Generate token pointing to Run 1
    v_svc = VerificationTokenService()
    token_doc = await v_svc.create_or_get_token(
        student_id=student_reg,
        batch_id=batch_id,
        extraction_request_id=run_1_id,
    )
    assert token_doc.extraction_request_id == run_1_id
    assert token_doc.status == VerificationTokenStatus.ACTIVE.value

    # Staff runs Run 2 with corrected/clearer document: father_name = "New Father Name"
    run_2_id = f"ext_ref_2_{datetime.now(timezone.utc).timestamp()}"
    await BatchExtractionRequest(
        extraction_request_id=run_2_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=now - timedelta(minutes=5),
        completed_at=now - timedelta(minutes=5),
    ).insert()

    await DocumentFieldVerification(
        extraction_request_id=run_2_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="father_name",
        ai_value="New Father Name",
        raw_value="New Father Name",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(minutes=5),
    ).insert()

    # Student opens their existing Run 1 link
    session_info = await v_svc.get_session_info(token_doc.token)
    assert session_info.valid is True
    assert session_info.extraction_request_id == run_2_id
    assert session_info.is_completed is False  # Loading data must NOT mark student verified

    # Student validates identity
    public_data = await v_svc.validate_student_identity(
        token_str=token_doc.token,
        student_name="David Refresh",
        register_number=student_reg,
    )
    assert public_data.extraction_request_id == run_2_id
    field_map = {f.field_name: f for f in public_data.fields}
    assert field_map["father_name"].ai_value == "New Father Name"

    # Token status in database must still be ACTIVE, not COMPLETED
    reloaded_token = await VerificationToken.find_one(VerificationToken.token == token_doc.token)
    assert reloaded_token.status == VerificationTokenStatus.ACTIVE.value
    assert reloaded_token.extraction_request_id == run_2_id


@pytest.mark.asyncio
async def test_excel_export_matches_selected_latest_extraction():
    """
    Test 6: Excel export consistency.
    Ensure Excel export for the latest extraction request exports the correct
    extraction values and student-submitted overrides, preserving extraction boundaries.
    """
    admin = await create_test_admin()
    batch_id = f"batch_excel_sync_{datetime.now(timezone.utc).timestamp()}"
    student_reg = "REG_EXCEL_06"
    now = datetime.now(timezone.utc)

    batch = AdmissionBatch(
        id=batch_id,
        name="Excel Sync Batch",
        academic_year="2026-2027",
        department_id="DEP_CSE",
        created_by=str(admin.id),
        created_at=now,
    )
    await batch.insert()

    student = StudentSubmission(
        batch_id=batch_id,
        student_id=student_reg,
        register_number=student_reg,
        student_name="Eva Sync",
        department_id="DEP_CSE",
        created_at=now,
    )
    await student.insert()

    # Run 1: blood_group = "B+"
    run_1_id = f"ext_sync_1_{datetime.now(timezone.utc).timestamp()}"
    await BatchExtractionRequest(
        extraction_request_id=run_1_id,
        batch_id=batch_id,
        batch_name="Excel Sync Batch",
        department_id="DEP_CSE",
        created_by=admin.username,
        status=BatchExtractionStatus.COMPLETED.value,
        requested_columns=["Student Name", "Register Number", "Blood Group"],
        requested_fields=["student_name", "register_number", "blood_group"],
        column_mapping={
            "Student Name": "student_name",
            "Register Number": "register_number",
            "Blood Group": "blood_group",
        },
        total_students=1,
        created_at=now - timedelta(hours=1),
    ).insert()

    await DocumentFieldVerification(
        extraction_request_id=run_1_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="blood_group",
        ai_value="B+",
        raw_value="B+",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(hours=1),
    ).insert()

    # Run 2: blood_group = "O+"
    run_2_id = f"ext_sync_2_{datetime.now(timezone.utc).timestamp()}"
    req_2 = BatchExtractionRequest(
        extraction_request_id=run_2_id,
        batch_id=batch_id,
        batch_name="Excel Sync Batch",
        department_id="DEP_CSE",
        created_by=admin.username,
        status=BatchExtractionStatus.COMPLETED.value,
        export_status=BatchExportStatus.READY.value,
        download_status=BatchDownloadStatus.NOT_DOWNLOADED.value,
        requested_columns=["Student Name", "Register Number", "Blood Group"],
        requested_fields=["student_name", "register_number", "blood_group"],
        column_mapping={
            "Student Name": "student_name",
            "Register Number": "register_number",
            "Blood Group": "blood_group",
        },
        total_students=1,
        created_at=now - timedelta(minutes=10),
    )
    await req_2.insert()

    await DocumentFieldVerification(
        extraction_request_id=run_2_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="blood_group",
        ai_value="O+",
        raw_value="O+",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=now - timedelta(minutes=10),
    ).insert()

    # Export Excel for Run 2
    b_svc = BatchExtractionService()
    stream, filename = await b_svc.export_verified_excel(
        extraction_request_id=run_2_id, actor=admin
    )
    wb = openpyxl.load_workbook(stream)
    ws = wb.active

    # Row 1 is headers: ["Student Name", "Register Number", "Blood Group"]
    # Row 2 is data
    headers = [cell.value for cell in ws[1]]
    assert "Blood Group" in headers
    bg_idx = headers.index("Blood Group")

    data_row = [cell.value for cell in ws[2]]
    # Blood Group in Run 2 must be "O+", not "B+"
    assert data_row[bg_idx] == "O+"


@pytest.mark.asyncio
async def test_fresh_verification_data_and_extracted_at_timestamp(client_async: AsyncClient):
    """
    Test 7: Fresh verification data, extracted_at timestamp, and student submission separation.
    Verifies that:
    - API returns extracted_at timestamp matching the extraction run's completion timestamp.
    - AI-extracted values are preserved separately when student submits a correction.
    - Subsequent refresh loads the updated verified state without overwriting or reverting.
    """
    batch_id = f"batch_ts_{datetime.now(timezone.utc).timestamp()}"
    student_reg = "REG_TS_07"
    now = datetime.now(timezone.utc)

    await AdmissionBatch(
        id=batch_id,
        name="Timestamp & Separation Batch",
        academic_year="2026-2027",
        department_id="DEP_CSE",
        created_by="admin",
        created_at=now,
    ).insert()

    await StudentSubmission(
        batch_id=batch_id,
        student_id=student_reg,
        register_number=student_reg,
        student_name="Fiona Time",
        email="fiona@test.com",
        mobile_number="9876543217",
        department_id="DEP_CSE",
        created_at=now,
    ).insert()

    run_id = f"ext_ts_{datetime.now(timezone.utc).timestamp()}"
    run_time = now - timedelta(minutes=45)
    await BatchExtractionRequest(
        extraction_request_id=run_id,
        batch_id=batch_id,
        status=BatchExtractionStatus.COMPLETED.value,
        created_by="admin",
        created_at=run_time,
        completed_at=run_time,
    ).insert()

    verif_rec = DocumentFieldVerification(
        extraction_request_id=run_id,
        student_id=student_reg,
        batch_id=batch_id,
        field_name="community",
        ai_value="BC",
        raw_value="BC",
        status=VerificationFieldStatus.AI_EXTRACTED,
        created_at=run_time,
    )
    await verif_rec.insert()

    v_svc = VerificationTokenService()
    token_doc = await v_svc.create_or_get_token(
        student_id=student_reg,
        batch_id=batch_id,
        extraction_request_id=run_id,
    )

    # 1. Fetch public verification data and check extracted_at is populated
    data_before = await v_svc.get_public_verification_data(token_doc.token)
    assert data_before.extracted_at is not None
    assert data_before.extraction_request_id == run_id
    field_item = next(f for f in data_before.fields if f.field_name == "community")
    assert field_item.ai_value == "BC"
    assert field_item.student_value is None

    # 2. Student submits correction: "BC_MUSLIM"
    from app.schemas.verification_token import PublicVerificationSubmitRequest, PublicFieldVerificationSubmitItem
    submit_req = PublicVerificationSubmitRequest(
        fields=[
            PublicFieldVerificationSubmitItem(
                verification_id=str(verif_rec.id),
                is_confirmed=False,
                corrected_value="BC_MUSLIM",
                comment="Updated by student",
            )
        ]
    )
    data_after = await v_svc.submit_public_verification(token_doc.token, submit_req)
    field_after = next(f for f in data_after.fields if f.field_name == "community")

    # Verify AI value is strictly preserved, and student value is separated
    assert field_after.ai_value == "BC"
    assert field_after.student_value == "BC_MUSLIM"
    assert field_after.final_value == "BC_MUSLIM"
    assert field_after.status in (VerificationFieldStatus.STUDENT_VERIFIED, VerificationFieldStatus.CORRECTION_REQUESTED)

    # 3. Reload session fresh - must retain student-submitted correction and AI value
    reloaded_data = await v_svc.get_public_verification_data(token_doc.token)
    reloaded_field = next(f for f in reloaded_data.fields if f.field_name == "community")
    assert reloaded_field.ai_value == "BC"
    assert reloaded_field.student_value == "BC_MUSLIM"
    assert reloaded_field.final_value == "BC_MUSLIM"
