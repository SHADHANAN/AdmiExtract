import io
import os
import openpyxl
import pytest
import pytest_asyncio
from datetime import datetime, timezone
from beanie import PydanticObjectId
from motor.motor_asyncio import AsyncIOMotorClient

from app.core.config import settings
from app.models.user import User, UserRole
from app.models.batch import AdmissionBatch
from app.models.student_submission import StudentSubmission, StudentDocumentMeta
from app.models.student_document import StudentDocument, DocumentVersion
from app.models.wanted_field_config import DocumentFieldConfiguration, WantedFieldItem
from app.models.document_verification import DocumentFieldVerification, VerificationFieldStatus
from app.models.batch_extraction import BatchExtractionRequest
from app.models.excel_template import ExcelBatchTemplate
from app.services.batch_extraction_service import BatchExtractionService
from app.services.excel_template_service import ExcelTemplateService


@pytest_asyncio.fixture(autouse=True)
async def setup_db():
    import app.db.database as db_mod

    db_mod.client = AsyncIOMotorClient(
        settings.MONGODB_URI, serverSelectionTimeoutMS=2000
    )
    db_mod.db = db_mod.client[settings.DATABASE_NAME]
    await db_mod._run_db_setup()


@pytest.mark.asyncio
async def test_extraction_configured_document_conflicts_aadhaar_vs_tc(tmp_path):
    """
    TEST 1 & TEST 2 (Requirements 23 & 24):
    Student has both Aadhaar and Transfer Certificate (TC).
    The documents contain conflicting values:
      - Aadhaar: Student Name = 'SHADHANAN', DOB = '12/05/2006'
      - TC: Student Name = 'SHADHANAN S', DOB = '13/05/2006'

    Scenario A:
      Configured: Student Name -> Aadhaar, DOB -> Aadhaar
      Expected: 'SHADHANAN', '12/05/2006'

    Scenario B:
      Configured: Student Name -> TC, DOB -> TC
      Expected: 'SHADHANAN S', '13/05/2006'
    """
    batch_service = BatchExtractionService()
    dept_id = f"DEPT_{PydanticObjectId()}"
    batch_id_a = f"batch_a_{PydanticObjectId()}"
    batch_id_b = f"batch_b_{PydanticObjectId()}"

    batch_a = AdmissionBatch(
        id=batch_id_a,
        name="Batch_Aadhaar_Config",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch_a.insert()

    batch_b = AdmissionBatch(
        id=batch_id_b,
        name="Batch_TC_Config",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch_b.insert()

    actor = User(
        name="Dept Admin",
        username=f"admin_{PydanticObjectId()}",
        email="admin@test.com",
        password="hash",
        role=UserRole.DEPARTMENT_ADMIN,
        department_code=dept_id,
        department_id=dept_id,
        is_active=True,
    )
    await actor.insert()

    # Create dummy files
    aadhaar_file = tmp_path / "aadhaar_shadhanan.pdf"
    aadhaar_file.write_text("AADHAAR CARD\nName: SHADHANAN\nDOB: 12/05/2006\nGender: MALE\n1234 5678 9012")

    tc_file = tmp_path / "tc_shadhanan.pdf"
    tc_file.write_text("TRANSFER CERTIFICATE\nName: SHADHANAN S\nDOB: 13/05/2006\nFather: SUBRAMANIAN")

    # Student 1 in Batch A
    stud1_aadhaar = StudentDocument(
        student_id="24AM093",
        batch_id=batch_id_a,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        versions=[
            DocumentVersion(
                version_id="v_aadhaar_1",
                version_number=1,
                file_path=str(aadhaar_file),
                original_filename="aadhaar.pdf",
                is_active=True,
            )
        ],
    )
    await stud1_aadhaar.insert()

    stud1_tc = StudentDocument(
        student_id="24AM093",
        batch_id=batch_id_a,
        document_type="TRANSFER_CERTIFICATE",
        document_name="Transfer Certificate",
        versions=[
            DocumentVersion(
                version_id="v_tc_1",
                version_number=1,
                file_path=str(tc_file),
                original_filename="tc.pdf",
                is_active=True,
            )
        ],
    )
    await stud1_tc.insert()

    sub_a = StudentSubmission(
        student_id="24AM093",
        register_number="24AM093",
        student_name="SHADHANAN S",
        batch_id=batch_id_a,
        department_id=dept_id,
        mobile_number="9876543210",
    )
    await sub_a.insert()

    # Configure Batch A: Fields configured under AADHAAR
    cfg_a = DocumentFieldConfiguration(
        batch_id=batch_id_a,
        document_type="AADHAAR",
        display_name="Aadhaar Card",
        requirement_status="REQUIRED",
        fields=[
            WantedFieldItem(field="Student Name", enabled=True, excel_header="Candidate Name"),
            WantedFieldItem(field="Date of Birth", enabled=True, excel_header="DOB"),
        ],
    )
    await cfg_a.insert()

    # Mock extract to return document-specific values
    async def mock_extract(doc, version, field_key, actor):
        from app.services.accuracy_extraction_engine import FieldExtractionResult
        norm_type = doc.document_type
        if norm_type == "AADHAAR":
            if field_key == "student_name":
                return FieldExtractionResult("student_name", "SHADHANAN", "SHADHANAN", "HIGH", 0.98, "EXTRACTED", "Name: SHADHANAN", source_document_type="AADHAAR")
            elif field_key in ("date_of_birth", "dob"):
                return FieldExtractionResult("date_of_birth", "12/05/2006", "12/05/2006", "HIGH", 0.98, "EXTRACTED", "DOB: 12/05/2006", source_document_type="AADHAAR")
        elif norm_type == "TRANSFER_CERTIFICATE":
            if field_key == "student_name":
                return FieldExtractionResult("student_name", "SHADHANAN S", "SHADHANAN S", "HIGH", 0.98, "EXTRACTED", "Name: SHADHANAN S", source_document_type="TRANSFER_CERTIFICATE")
            elif field_key in ("date_of_birth", "dob"):
                return FieldExtractionResult("date_of_birth", "13/05/2006", "13/05/2006", "HIGH", 0.98, "EXTRACTED", "DOB: 13/05/2006", source_document_type="TRANSFER_CERTIFICATE")
        return FieldExtractionResult(field_key, None, "NOT_AVAILABLE", "LOW", 0.0, "NOT_FOUND", source_document_type=norm_type)

    from unittest.mock import patch
    with patch.object(BatchExtractionService, "_extract_field_with_accuracy", side_effect=mock_extract):
        # Run Extraction for Batch A
        res_a = await batch_service.create_and_run_batch_extraction(
            batch_id=batch_id_a,
            requested_columns=["Candidate Name", "DOB"],
            requested_fields=["student_name", "date_of_birth"],
            actor=actor,
        )

    # Verify Batch A extracted from Aadhaar ONLY
    details_a = await batch_service.get_extraction_details(res_a.extraction_request_id, actor)
    assert len(details_a.students) == 1
    stud_a = details_a.students[0]
    assert stud_a.fields["student_name"].ai_value == "SHADHANAN"
    assert stud_a.fields["student_name"].source_document_type == "AADHAAR"
    assert stud_a.fields["date_of_birth"].ai_value == "12/05/2006"
    assert stud_a.fields["date_of_birth"].source_document_type == "AADHAAR"

    # Now Setup Batch B: Configured under TRANSFER_CERTIFICATE
    stud2_aadhaar = StudentDocument(
        student_id="24AM094",
        batch_id=batch_id_b,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        versions=[
            DocumentVersion(
                version_id="v_aadhaar_2",
                version_number=1,
                file_path=str(aadhaar_file),
                original_filename="aadhaar.pdf",
                is_active=True,
            )
        ],
    )
    await stud2_aadhaar.insert()

    stud2_tc = StudentDocument(
        student_id="24AM094",
        batch_id=batch_id_b,
        document_type="TRANSFER_CERTIFICATE",
        document_name="Transfer Certificate",
        versions=[
            DocumentVersion(
                version_id="v_tc_2",
                version_number=1,
                file_path=str(tc_file),
                original_filename="tc.pdf",
                is_active=True,
            )
        ],
    )
    await stud2_tc.insert()

    sub_b = StudentSubmission(
        student_id="24AM094",
        register_number="24AM094",
        student_name="SHADHANAN",
        batch_id=batch_id_b,
        department_id=dept_id,
        mobile_number="9876543210",
    )
    await sub_b.insert()

    cfg_b = DocumentFieldConfiguration(
        batch_id=batch_id_b,
        document_type="TRANSFER_CERTIFICATE",
        display_name="Transfer Certificate",
        requirement_status="REQUIRED",
        fields=[
            WantedFieldItem(field="Student Name", enabled=True, excel_header="Candidate Name"),
            WantedFieldItem(field="Date of Birth", enabled=True, excel_header="DOB"),
        ],
    )
    await cfg_b.insert()

    with patch.object(BatchExtractionService, "_extract_field_with_accuracy", side_effect=mock_extract):
        # Run Extraction for Batch B
        res_b = await batch_service.create_and_run_batch_extraction(
            batch_id=batch_id_b,
            requested_columns=["Candidate Name", "DOB"],
            requested_fields=["student_name", "date_of_birth"],
            actor=actor,
        )

    # Verify Batch B extracted from TC ONLY
    details_b = await batch_service.get_extraction_details(res_b.extraction_request_id, actor)
    assert len(details_b.students) == 1
    stud_b = details_b.students[0]
    assert stud_b.fields["student_name"].ai_value == "SHADHANAN S"
    assert stud_b.fields["student_name"].source_document_type == "TRANSFER_CERTIFICATE"
    assert stud_b.fields["date_of_birth"].ai_value == "13/05/2006"
    assert stud_b.fields["date_of_birth"].source_document_type == "TRANSFER_CERTIFICATE"


@pytest.mark.asyncio
async def test_extraction_reupload_active_version_rule(tmp_path):
    """
    TEST 4 (Requirement 26):
    Document reupload versioning rule:
      Aadhaar V1 = 'OLD STUDENT NAME' (is_active=False)
      Aadhaar V2 = 'NEW STUDENT NAME' (is_active=True)
    Expected:
      Extraction engine extracts from V2 ONLY. Never V1.
    """
    batch_service = BatchExtractionService()
    dept_id = f"DEPT_{PydanticObjectId()}"
    batch_id = f"batch_reupload_{PydanticObjectId()}"

    batch = AdmissionBatch(
        id=batch_id,
        name="Batch_Reupload_Test",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch.insert()

    actor = User(
        name="Dept Admin",
        username=f"admin_{PydanticObjectId()}",
        email="admin@test.com",
        password="hash",
        role=UserRole.DEPARTMENT_ADMIN,
        department_code=dept_id,
        department_id=dept_id,
        is_active=True,
    )
    await actor.insert()

    v1_file = tmp_path / "aadhaar_v1.pdf"
    v1_file.write_text("AADHAAR V1\nName: OLD STUDENT NAME")

    v2_file = tmp_path / "aadhaar_v2.pdf"
    v2_file.write_text("AADHAAR V2\nName: NEW STUDENT NAME")

    doc = StudentDocument(
        student_id="24AM095",
        batch_id=batch_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        current_version=2,
        active_version_id="v_aadhaar_2",
        versions=[
            DocumentVersion(
                version_id="v_aadhaar_1",
                version_number=1,
                file_path=str(v1_file),
                original_filename="aadhaar_old.pdf",
                is_active=False,
            ),
            DocumentVersion(
                version_id="v_aadhaar_2",
                version_number=2,
                file_path=str(v2_file),
                original_filename="aadhaar_new.pdf",
                is_active=True,
            ),
        ],
    )
    await doc.insert()

    sub = StudentSubmission(
        student_id="24AM095",
        register_number="24AM095",
        student_name="NEW STUDENT NAME",
        batch_id=batch_id,
        department_id=dept_id,
        mobile_number="9876543210",
    )
    await sub.insert()

    cfg = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="AADHAAR",
        display_name="Aadhaar Card",
        requirement_status="REQUIRED",
        fields=[
            WantedFieldItem(field="Student Name", enabled=True, excel_header="Student Name"),
        ],
    )
    await cfg.insert()

    async def mock_extract_version(doc, version, field_key, actor):
        from app.services.accuracy_extraction_engine import FieldExtractionResult
        if version.version_number == 1:
            return FieldExtractionResult(field_key, "OLD STUDENT NAME", "OLD STUDENT NAME", "HIGH", 0.9, "EXTRACTED", source_document_type="AADHAAR")
        elif version.version_number == 2:
            return FieldExtractionResult(field_key, "NEW STUDENT NAME", "NEW STUDENT NAME", "HIGH", 0.99, "EXTRACTED", source_document_type="AADHAAR")
        return FieldExtractionResult(field_key, None, "NOT_AVAILABLE", "LOW", 0.0, "NOT_FOUND", source_document_type="AADHAAR")

    from unittest.mock import patch
    with patch.object(BatchExtractionService, "_extract_field_with_accuracy", side_effect=mock_extract_version):
        res = await batch_service.create_and_run_batch_extraction(
            batch_id=batch_id,
            requested_columns=["Student Name"],
            requested_fields=["student_name"],
            actor=actor,
        )

    details = await batch_service.get_extraction_details(res.extraction_request_id, actor)
    assert len(details.students) == 1
    stud = details.students[0]
    assert stud.fields["student_name"].ai_value == "NEW STUDENT NAME"
    assert stud.fields["student_name"].source_document_type == "AADHAAR"


@pytest.mark.asyncio
async def test_extraction_missing_document_handling(tmp_path):
    """
    TEST 5 (Requirement 27):
    Configuration:
      - Aadhaar -> Student Name
      - Bank Passbook -> Account Number
      - Community Certificate -> Community
    Student has Aadhaar and Bank Passbook, but NO Community Certificate.
    Expected:
      - Student Name: EXTRACTED from Aadhaar
      - Account Number: EXTRACTED from Bank Passbook
      - Community: NOT_AVAILABLE with status SOURCE_DOCUMENT_MISSING
      - Missing document does NOT fail the whole student or other fields.
    """
    batch_service = BatchExtractionService()
    dept_id = f"DEPT_{PydanticObjectId()}"
    batch_id = f"batch_missing_doc_{PydanticObjectId()}"

    batch = AdmissionBatch(
        id=batch_id,
        name="Batch_Missing_Doc_Test",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch.insert()

    actor = User(
        name="Dept Admin",
        username=f"admin_{PydanticObjectId()}",
        email="admin@test.com",
        password="hash",
        role=UserRole.DEPARTMENT_ADMIN,
        department_code=dept_id,
        department_id=dept_id,
        is_active=True,
    )
    await actor.insert()

    aadhaar_file = tmp_path / "aadhaar_doc.pdf"
    aadhaar_file.write_text("AADHAAR CARD\nName: PRAVEEN")

    bank_file = tmp_path / "passbook_doc.pdf"
    bank_file.write_text("BANK PASSBOOK\nAccount: 987654321098")

    # Stored docs: ONLY Aadhaar and Bank
    doc_aadhaar = StudentDocument(
        student_id="24AM096",
        batch_id=batch_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        versions=[
            DocumentVersion(
                version_id="v_a_1",
                version_number=1,
                file_path=str(aadhaar_file),
                original_filename="aadhaar.pdf",
                is_active=True,
            )
        ],
    )
    await doc_aadhaar.insert()

    doc_bank = StudentDocument(
        student_id="24AM096",
        batch_id=batch_id,
        document_type="BANK_PASSBOOK",
        document_name="Bank Passbook",
        versions=[
            DocumentVersion(
                version_id="v_b_1",
                version_number=1,
                file_path=str(bank_file),
                original_filename="bank.pdf",
                is_active=True,
            )
        ],
    )
    await doc_bank.insert()

    sub = StudentSubmission(
        student_id="24AM096",
        register_number="24AM096",
        student_name="PRAVEEN",
        batch_id=batch_id,
        department_id=dept_id,
        mobile_number="9876543210",
    )
    await sub.insert()

    # Configurations for 3 documents
    cfg_aadhaar = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="AADHAAR",
        display_name="Aadhaar Card",
        fields=[WantedFieldItem(field="Student Name", enabled=True, excel_header="Student Name")],
    )
    await cfg_aadhaar.insert()

    cfg_bank = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="BANK_PASSBOOK",
        display_name="Bank Passbook",
        fields=[WantedFieldItem(field="Account Number", enabled=True, excel_header="Account No")],
    )
    await cfg_bank.insert()

    cfg_comm = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="COMMUNITY",
        display_name="Community Certificate",
        fields=[WantedFieldItem(field="Community", enabled=True, excel_header="Community Category")],
    )
    await cfg_comm.insert()

    async def mock_extract_multi(doc, version, field_key, actor):
        from app.services.accuracy_extraction_engine import FieldExtractionResult
        if doc.document_type == "AADHAAR" and field_key == "student_name":
            return FieldExtractionResult("student_name", "PRAVEEN", "PRAVEEN", "HIGH", 0.98, "EXTRACTED", source_document_type="AADHAAR")
        elif doc.document_type == "BANK_PASSBOOK" and field_key == "account_number":
            return FieldExtractionResult("account_number", "987654321098", "987654321098", "HIGH", 0.98, "EXTRACTED", source_document_type="BANK_PASSBOOK")
        return FieldExtractionResult(field_key, None, "NOT_AVAILABLE", "LOW", 0.0, "NOT_FOUND", source_document_type=doc.document_type)

    from unittest.mock import patch
    with patch.object(BatchExtractionService, "_extract_field_with_accuracy", side_effect=mock_extract_multi):
        res = await batch_service.create_and_run_batch_extraction(
            batch_id=batch_id,
            requested_columns=["Student Name", "Account No", "Community Category"],
            requested_fields=["student_name", "account_number", "community_category"],
            actor=actor,
        )

    assert res.partial_count == 1
    assert res.missing_docs_count == 1

    details = await batch_service.get_extraction_details(res.extraction_request_id, actor)
    stud = details.students[0]
    assert stud.fields["student_name"].ai_value == "PRAVEEN"
    assert stud.fields["account_number"].ai_value == "987654321098"
    assert stud.fields["community_category"].ai_value == "NOT_AVAILABLE"
    assert stud.fields["community_category"].field_state == "SOURCE_DOCUMENT_MISSING"
    assert stud.fields["community_category"].source_document_type == "COMMUNITY"


@pytest.mark.asyncio
async def test_student_extraction_isolation(tmp_path):
    """
    TEST 6 (Requirement 28):
    Student A (Aadhaar Name = 'STUDENT AAA')
    Student B (Aadhaar Name = 'STUDENT BBB')
    Extraction context is strictly isolated per student.
    """
    batch_service = BatchExtractionService()
    dept_id = f"DEPT_{PydanticObjectId()}"
    batch_id = f"batch_iso_{PydanticObjectId()}"

    batch = AdmissionBatch(
        id=batch_id,
        name="Batch_Isolation_Test",
        academic_year="2025-2026",
        department_id=dept_id,
        created_by="admin_user",
        created_at=datetime.now(timezone.utc),
    )
    await batch.insert()

    actor = User(
        name="Dept Admin",
        username=f"admin_{PydanticObjectId()}",
        email="admin@test.com",
        password="hash",
        role=UserRole.DEPARTMENT_ADMIN,
        department_code=dept_id,
        department_id=dept_id,
        is_active=True,
    )
    await actor.insert()

    f_a = tmp_path / "a_aadhaar.pdf"
    f_a.write_text("AADHAAR\nName: STUDENT AAA")

    f_b = tmp_path / "b_aadhaar.pdf"
    f_b.write_text("AADHAAR\nName: STUDENT BBB")

    doc_a = StudentDocument(
        student_id="24AM101",
        batch_id=batch_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        versions=[DocumentVersion(version_id="v_a", version_number=1, file_path=str(f_a), original_filename="a.pdf", is_active=True)],
    )
    await doc_a.insert()

    doc_b = StudentDocument(
        student_id="24AM102",
        batch_id=batch_id,
        document_type="AADHAAR",
        document_name="Aadhaar Card",
        versions=[DocumentVersion(version_id="v_b", version_number=1, file_path=str(f_b), original_filename="b.pdf", is_active=True)],
    )
    await doc_b.insert()

    sub_a = StudentSubmission(student_id="24AM101", register_number="24AM101", student_name="STUDENT AAA", mobile_number="9876543210", batch_id=batch_id, department_id=dept_id)
    await sub_a.insert()

    sub_b = StudentSubmission(student_id="24AM102", register_number="24AM102", student_name="STUDENT BBB", mobile_number="9876543210", batch_id=batch_id, department_id=dept_id)
    await sub_b.insert()

    cfg = DocumentFieldConfiguration(
        batch_id=batch_id,
        document_type="AADHAAR",
        display_name="Aadhaar Card",
        fields=[WantedFieldItem(field="Student Name", enabled=True, excel_header="Student Name")],
    )
    await cfg.insert()

    async def mock_extract_iso(doc, version, field_key, actor):
        from app.services.accuracy_extraction_engine import FieldExtractionResult
        if doc.student_id == "24AM101":
            return FieldExtractionResult(field_key, "STUDENT AAA", "STUDENT AAA", "HIGH", 0.98, "EXTRACTED", source_document_type="AADHAAR")
        elif doc.student_id == "24AM102":
            return FieldExtractionResult(field_key, "STUDENT BBB", "STUDENT BBB", "HIGH", 0.98, "EXTRACTED", source_document_type="AADHAAR")
        return FieldExtractionResult(field_key, None, "NOT_AVAILABLE", "LOW", 0.0, "NOT_FOUND", source_document_type="AADHAAR")

    from unittest.mock import patch
    with patch.object(BatchExtractionService, "_extract_field_with_accuracy", side_effect=mock_extract_iso):
        res = await batch_service.create_and_run_batch_extraction(
            batch_id=batch_id,
            requested_columns=["Student Name"],
            requested_fields=["student_name"],
            actor=actor,
        )

    details = await batch_service.get_extraction_details(res.extraction_request_id, actor)
    assert len(details.students) == 2
    stud_map = {s.student_id: s for s in details.students}
    assert stud_map["24AM101"].fields["student_name"].ai_value == "STUDENT AAA"
    assert stud_map["24AM102"].fields["student_name"].ai_value == "STUDENT BBB"
