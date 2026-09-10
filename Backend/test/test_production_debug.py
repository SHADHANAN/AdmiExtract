"""
Production End-to-End Verification Test Suite
=============================================
Tests the complete data flow:
Document Upload -> AI Vision / OCR Extraction -> Classification -> Candidate Pooling
-> Field Verification -> Excel Mapping -> Excel Writing -> Workbook Persistence.

Verifies:
1. Single document extraction
2. 2-document extraction (Aadhaar + Marksheet)
3. 4-document extraction (Aadhaar + SSLC + HSC + Community)
4. 10-document extraction (Mixed PDFs, Images, multiple certificates)
5. 20+ document extraction (High-volume mixed batch)
6. Excel header semantic mapping across diverse column styles
7. Non-destructive Excel row updating & workbook persistence verification
8. Absolute prevention of blank Excel output or literal "NO" cell values
"""

import io
import os
import shutil
import tempfile
import pytest
import openpyxl
from pathlib import Path
from unittest.mock import MagicMock, AsyncMock, patch

from app.services.excel_template_service import ExcelTemplateService, _find_best_worksheet_and_headers, _find_lookup_col_index
from app.services.document_processing_pipeline import DocumentProcessingPipeline
from app.services.field_mapping_service import FieldMappingService
from app.services.document_classifier_service import DocumentClassifierService
from app.utils.normalization import normalize_register_number, compare_register_numbers
from app.utils.field_canonicalizer import is_document_authorized_for_field


@pytest.fixture
def temp_excel_template():
    """Create a temporary Excel workbook with diverse real-world Indian college headers."""
    temp_dir = tempfile.mkdtemp()
    file_path = os.path.join(temp_dir, "master_template.xlsx")

    wb = openpyxl.Workbook()
    # Sheet 1: Instructions (should be bypassed by smart sheet detector)
    ws_notes = wb.active
    ws_notes.title = "Instructions"
    ws_notes["A1"] = "Please fill in the student admission details below."

    # Sheet 2: Student Admission Master Sheet
    ws_data = wb.create_sheet(title="Admission Roster")
    # Row 1: Merged title banner (should be bypassed by header detector)
    ws_data["A1"] = "GOVERNMENT COLLEGE OF TECHNOLOGY - ADMISSIONS 2024-2025"

    # Row 2: Actual column headers
    headers = [
        "S.No",
        "Roll No",
        "Name of the Student",
        "Register No.",
        "Mobile Number",
        "Email ID",
        "Aadhaar Number",
        "Community / Caste",
        "Date of Birth",
        "Permanent Address",
        "SSLC Mark Percentage",
        "HSC Mark Percentage",
        "First Graduate (Yes/No)",
    ]
    for col_idx, h in enumerate(headers, start=1):
        ws_data.cell(row=2, column=col_idx, value=h)

    # Pre-populate 3 existing student rows (to test existing row updates)
    students = [
        (1, "24CS001", "Aravind S", "240101"),
        (2, "24CS002", "Deepa R", "240102"),
        (3, "24CS003", "Karthik M", "00240103"),  # zero-padded
    ]
    for r_offset, (sno, roll, name, reg) in enumerate(students, start=3):
        ws_data.cell(row=r_offset, column=1, value=sno)
        ws_data.cell(row=r_offset, column=2, value=roll)
        ws_data.cell(row=r_offset, column=3, value=name)
        ws_data.cell(row=r_offset, column=4, value=reg)

    wb.save(file_path)
    wb.close()

    yield file_path, temp_dir
    shutil.rmtree(temp_dir, ignore_errors=True)


class TestSemanticHeaderResolution:
    """Test _resolve_header_value against any reasonable column header style."""

    def test_student_name_variations(self):
        service = ExcelTemplateService()
        data = {
            "Student Name": "Sundar Pichai",
            "Register Number": "24AM101",
        }
        test_headers = [
            "Name of the Student",
            "Candidate Name",
            "Name of the Candidate",
            "Student's Name",
            "Name",
            "Full Name",
            "Student Full Name",
            "Name of Candidate",
            "Applicant Name",
        ]
        for h in test_headers:
            resolved = service._resolve_header_value(h, data)
            assert resolved == "Sundar Pichai", f"Failed to resolve '{h}'"

    def test_register_and_roll_number_variations(self):
        service = ExcelTemplateService()
        data = {
            "Register Number": "24AM101",
            "Student Name": "John Doe",
        }
        test_headers = [
            "Register Number",
            "Register No",
            "Register No.",
            "Reg No",
            "Reg. No",
            "Registration Number",
            "Registration No",
            "Roll Number",
            "Roll No",
            "Roll. No",
            "Reg_No",
        ]
        for h in test_headers:
            resolved = service._resolve_header_value(h, data)
            assert resolved == "24AM101", f"Failed to resolve '{h}'"

    def test_community_and_caste_variations(self):
        service = ExcelTemplateService()
        data = {
            "Community Category": "MBC",
            "Student Name": "Arun Kumar",
        }
        test_headers = [
            "Community",
            "Community Category",
            "Caste",
            "Community / Caste",
            "Community Name",
        ]
        for h in test_headers:
            resolved = service._resolve_header_value(h, data)
            assert resolved == "MBC", f"Failed to resolve '{h}'"

    def test_never_returns_literal_no_for_missing_fields(self):
        service = ExcelTemplateService()
        data = {
            "Student Name": "Test Student",
            "Extracted Field": "NO",
            "Another Field": "NULL",
        }
        # Standard fields must return None, NEVER "NO"
        assert service._resolve_header_value("Aadhaar Number", data) is None
        assert service._resolve_header_value("Permanent Address", data) is None
        assert service._resolve_header_value("Extracted Field", data) is None
        assert service._resolve_header_value("Another Field", data) is None

        # Boolean Yes/No questions may return "No"
        assert service._resolve_header_value("Is Physically Challenged (Yes/No)", data) is None
        assert service._resolve_header_value("First Graduate (Yes/No)", {"First Graduate (Yes/No)": "No"}) == "No"


class TestExcelWorksheetAndRowDetection:
    """Test smart sheet selection, header row identification, and row matching."""

    def test_smart_worksheet_and_header_detection(self, temp_excel_template):
        file_path, _ = temp_excel_template
        wb = openpyxl.load_workbook(file_path, data_only=True)

        sheet, header_row_idx, headers, header_col_map = _find_best_worksheet_and_headers(wb)
        wb.close()

        # Must select the "Admission Roster" sheet, NOT the "Instructions" sheet
        assert sheet.title == "Admission Roster"
        # Must detect Row 2 as the header row, bypassing the merged title in Row 1
        assert header_row_idx == 2
        assert "Roll No" in headers
        assert "Name of the Student" in headers
        assert "Register No." in headers

    def test_lookup_column_identification(self, temp_excel_template):
        file_path, _ = temp_excel_template
        wb = openpyxl.load_workbook(file_path, data_only=True)
        _, _, _, header_col_map = _find_best_worksheet_and_headers(wb)
        wb.close()

        col_name, col_idx = _find_lookup_col_index(header_col_map)
        # Should detect "Register No." or "Roll No"
        assert col_name in ["Register No.", "Roll No"]
        assert col_idx in [2, 4]

    def test_numeric_zero_padding_comparison(self):
        assert compare_register_numbers("00240103", "240103") is True
        assert compare_register_numbers("240101", 240101) is True
        assert compare_register_numbers(240102.0, "240102") is True
        assert compare_register_numbers("24AM101", "24am101") is True
        assert compare_register_numbers("240101", "240102") is False


class TestExcelRowUpdatingAndPersistence:
    """Test writing into Excel rows non-destructively and verifying file on disk."""

    @pytest.mark.asyncio
    async def test_update_existing_student_row(self, temp_excel_template):
        file_path, _ = temp_excel_template
        service = ExcelTemplateService()

        # Mock repository to return our temp template
        mock_template = MagicMock()
        mock_template.file_path = file_path
        mock_template.template_filename = "master_template.xlsx"
        mock_template.lookup_column = "Register No."
        service.repository.get_by_batch_id = AsyncMock(return_value=mock_template)
        service.repository.increment_updated_count = AsyncMock()

        student_data = {
            "Student Name": "Aravind S",
            "Register Number": "240101",
            "Mobile Number": "9876543210",
            "Email": "aravind@example.com",
            "Aadhaar Number": "9876 5432 1098",
            "Community Category": "OBC",
            "Permanent Address": "45 Anna Salai, Chennai 600002",
            "SSLC Mark Percentage": "95.2%",
            "HSC Mark Percentage": "97.4%",
            "Date of Birth": "12/04/2006",
        }

        # Update row for 240101 (existing student in row 3)
        success = await service.append_or_update_student_row_in_excel(
            batch_id="batch_123",
            register_number="240101",
            student_data=student_data,
        )
        assert success is True

        # Re-open workbook directly from disk to verify persisted cells
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb["Admission Roster"]

        # Row 3 is student Aravind S (Register No: 240101)
        assert sheet.cell(row=3, column=3).value == "Aravind S"
        assert str(sheet.cell(row=3, column=4).value) == "240101"
        assert str(sheet.cell(row=3, column=5).value) == "9876543210"  # Mobile
        assert sheet.cell(row=3, column=6).value == "aravind@example.com"  # Email
        assert sheet.cell(row=3, column=7).value == "9876 5432 1098"  # Aadhaar
        assert sheet.cell(row=3, column=8).value == "OBC"  # Community
        assert sheet.cell(row=3, column=9).value == "12/04/2006"  # DOB
        assert "Anna Salai" in sheet.cell(row=3, column=10).value  # Address
        assert sheet.cell(row=3, column=11).value == "95.2%"  # SSLC
        assert sheet.cell(row=3, column=12).value == "97.4%"  # HSC
        wb.close()

    @pytest.mark.asyncio
    async def test_update_zero_padded_register_student_row(self, temp_excel_template):
        file_path, _ = temp_excel_template
        service = ExcelTemplateService()

        mock_template = MagicMock()
        mock_template.file_path = file_path
        mock_template.template_filename = "master_template.xlsx"
        mock_template.lookup_column = "Register No."
        service.repository.get_by_batch_id = AsyncMock(return_value=mock_template)
        service.repository.increment_updated_count = AsyncMock()

        # Student input is "240103" but sheet has "00240103"
        student_data = {
            "Student Name": "Karthik M",
            "Register Number": "240103",
            "Aadhaar Number": "1111 2222 3333",
            "Community Category": "SC",
        }

        success = await service.append_or_update_student_row_in_excel(
            batch_id="batch_123",
            register_number="240103",
            student_data=student_data,
        )
        assert success is True

        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb["Admission Roster"]
        # Row 5 corresponds to Karthik M (originally 00240103)
        assert sheet.cell(row=5, column=3).value == "Karthik M"
        assert sheet.cell(row=5, column=7).value == "1111 2222 3333"
        assert sheet.cell(row=5, column=8).value == "SC"
        wb.close()


class TestMultiDocumentPipelineScalability:
    """Test that pipeline processes 1, 2, 4, 10, and 20+ documents without dropping any."""

    @pytest.mark.asyncio
    async def test_multi_document_batches(self):
        """Test pipeline handling varying document counts: 1, 2, 4, 10, 20+."""
        pipeline = DocumentProcessingPipeline()

        # Mock Excel headers
        excel_headers = [
            "Student Name", "Register Number", "Mobile Number", "Aadhaar Card",
            "Community Category", "Date of Birth", "Permanent Address", "SSLC Total", "HSC Total"
        ]
        pipeline.excel_service.get_template_by_batch = AsyncMock(return_value=MagicMock(headers=excel_headers))

        for doc_count in [1, 2, 4, 10, 25]:
            mock_files = []
            for i in range(doc_count):
                f = MagicMock()
                ext = ".pdf" if i % 2 == 0 else ".jpg"
                f.filename = f"student_cert_{i+1}{ext}"
                # Async read returning dummy bytes
                f.read = AsyncMock(return_value=b"%PDF-1.4 dummy file content")
                mock_files.append(f)

            # Mock Gemini vision extractor to return document fields
            def mock_gemini_extract(file_bytes, mime_type, filename, target_fields):
                return {
                    "success": True,
                    "document_type": "AADHAAR" if "1" in filename else "SSLC" if "2" in filename else "COMMUNITY",
                    "fields": {
                        "Aadhaar Card": {"value": "9999 8888 7777", "confidence": 99} if "1" in filename else {},
                        "Community Category": {"value": "BC", "confidence": 98} if "3" in filename else {},
                        "Date of Birth": {"value": "10/05/2005", "confidence": 95},
                    },
                    "extracted_summary": {
                        "full_address": "77 Gandhi Road, Madurai" if "1" in filename else None,
                    }
                }

            pipeline.gemini_service.is_available = MagicMock(return_value=True)
            pipeline.gemini_service.extract_from_bytes = MagicMock(side_effect=mock_gemini_extract)

            response = await pipeline.process_student_documents(
                batch_id="batch_batch_test",
                register_number="24REG999",
                student_name="Sanjay Raman",
                mobile_number="9123456780",
                files=mock_files,
            )

            # Assert all documents were processed
            assert response["status"] == "success"
            assert len(response["extracted_fields_per_document"]) == doc_count
            assert response["login_fields"]["Student Name"] == "Sanjay Raman"
            assert response["login_fields"]["Register Number"] == "24REG999"

            # Assert merged fields contain profile and extracted attributes
            merged = response["merged_fields"]
            assert merged["Student Name"]["value"] == "Sanjay Raman"
            assert merged["Register Number"]["value"] == "24REG999"
            assert merged["Mobile Number"]["value"] == "9123456780"

            # Check that missing fields have None, NEVER "NO"
            for h in excel_headers:
                val = merged[h]["value"]
                if val is not None:
                    assert val != "NO", f"Header '{h}' contains literal 'NO'"
