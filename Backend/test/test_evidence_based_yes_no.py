"""
Document-Evidence-Based YES/NO Answers Test Suite
=================================================
Validates strict document-evidence rules for YES / NO / NOT_FOUND:
1. Explicit YES in document -> YES
2. Explicit NO in document -> NO
3. Document contains sufficient positive evidence -> YES
4. Document contains sufficient negative evidence -> NO
5. Document does not contain enough evidence -> NOT_FOUND
6. Never guess.
7. Never infer NO merely because a field is missing.
8. Never fabricate an answer.
9. NOT_FOUND must NEVER be converted to NO.
10. Derived questions (communication_address_same) evaluated by address comparison.
11. Answer dictionary contains field, value, confidence, evidence, source_document, source_page.
12. Excel export writes strictly YES / NO / NOT_FOUND without shifting columns.
"""

import io
import openpyxl
import pytest

from app.utils.evidence_yes_no import (
    is_yes_no_field,
    normalize_yes_no_not_found,
    compare_addresses_for_similarity,
    evaluate_evidence_based_yes_no,
)
from app.services.accuracy_extraction_engine import (
    AccuracyExtractionEngine,
    FieldExtractionResult,
)
from app.services.field_mapping_service import (
    FieldMappingService,
    format_for_excel_write,
)


# =============================================================================
# 1. Explicit YES Tests
# =============================================================================

def test_explicit_yes_first_graduate():
    doc_text = "Transfer Certificate\nFirst Graduate: Yes\nCandidate Name: John Doe"
    ans = evaluate_evidence_based_yes_no(
        field_key="is_first_graduate",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
        page_number=1,
    )
    assert ans["field"] == "is_first_graduate"
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.9
    assert ans["evidence"] is not None
    assert "First Graduate: Yes" in ans["evidence"]
    assert ans["source_document"] == "TRANSFER_CERTIFICATE"
    assert ans["source_page"] == 1


def test_explicit_yes_differently_abled():
    doc_text = "Student Profile\nDifferently Abled: Yes (Visually Impaired 40%)"
    ans = evaluate_evidence_based_yes_no(
        field_key="differently_abled",
        document_text=doc_text,
        document_type="COMMUNITY",
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.9
    assert ans["evidence"] is not None
    assert "Differently Abled: Yes" in ans["evidence"]


def test_explicit_yes_orphan():
    doc_text = "Admission Application\nOrphan: Yes\nCategory: General"
    ans = evaluate_evidence_based_yes_no(
        field_key="orphan_category",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.9
    assert "Orphan: Yes" in ans["evidence"]


def test_explicit_yes_raw_value_normalized():
    # If raw extracted value from Gemini was affirmative
    ans = evaluate_evidence_based_yes_no(
        field_key="is_first_graduate",
        raw_val="Yes",
        document_text="",
        document_type="BONAFIDE",
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.9


# =============================================================================
# 2. Explicit NO Tests
# =============================================================================

def test_explicit_no_first_graduate():
    doc_text = "Transfer Certificate\nFirst Graduate: No\nCandidate Name: Priya"
    ans = evaluate_evidence_based_yes_no(
        field_key="is_first_graduate",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
        page_number=1,
    )
    assert ans["field"] == "is_first_graduate"
    assert ans["value"] == "NO"
    assert ans["confidence"] >= 0.9
    assert ans["evidence"] is not None
    assert "First Graduate: No" in ans["evidence"]


def test_explicit_no_differently_abled():
    doc_text = "Student Health Sheet\nDifferently Abled: No\nBlood Group: O+"
    ans = evaluate_evidence_based_yes_no(
        field_key="differently_abled",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "NO"
    assert ans["confidence"] >= 0.9
    assert "Differently Abled: No" in ans["evidence"]


def test_explicit_no_orphan():
    doc_text = "School Certificate\nOrphan: No\nFather Alive: Yes"
    ans = evaluate_evidence_based_yes_no(
        field_key="orphan_category",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "NO"
    assert ans["confidence"] >= 0.9
    assert "Orphan: No" in ans["evidence"]


def test_explicit_no_raw_value_normalized():
    ans = evaluate_evidence_based_yes_no(
        field_key="special_admission_quota",
        raw_val="No",
        document_text="",
        document_type="ALLOTMENT_ORDER",
    )
    assert ans["value"] == "NO"
    assert ans["confidence"] >= 0.9


# =============================================================================
# 3. Missing Information & Blank Field Tests (NEVER Convert to NO)
# =============================================================================

def test_missing_disability_returns_not_found_never_no():
    # Document contains general school info but no disability mention
    doc_text = "Secondary School Leaving Certificate\nName: Ramesh\nDOB: 12/05/2005"
    ans = evaluate_evidence_based_yes_no(
        field_key="differently_abled",
        document_text=doc_text,
        document_type="SSLC",
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["value"] != "NO"  # CRITICAL: NOT_FOUND must NEVER be converted to NO
    assert ans["confidence"] == 0.0
    assert ans["evidence"] is None


def test_missing_orphan_returns_not_found_never_no():
    doc_text = "Aadhaar Card\nGovernment of India\nFather: Suresh"
    ans = evaluate_evidence_based_yes_no(
        field_key="orphan_category",
        document_text=doc_text,
        document_type="AADHAAR",
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["value"] != "NO"
    assert ans["confidence"] == 0.0
    assert ans["evidence"] is None


def test_blank_field_in_document_returns_not_found():
    # "First Graduate: [blank]" or empty space
    doc_text = "Application Form\nFirst Graduate: [blank]\nReligion: Hindu"
    ans = evaluate_evidence_based_yes_no(
        field_key="is_first_graduate",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["value"] != "NO"
    assert ans["confidence"] == 0.0
    assert ans["evidence"] is None


def test_empty_string_raw_value_returns_not_found():
    ans = evaluate_evidence_based_yes_no(
        field_key="is_first_graduate",
        raw_val="",
        document_text="",
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["confidence"] == 0.0
    assert ans["evidence"] is None


def test_normalizer_never_coerces_not_found_to_no():
    assert normalize_yes_no_not_found(None) == "NOT_FOUND"
    assert normalize_yes_no_not_found("") == "NOT_FOUND"
    assert normalize_yes_no_not_found("   ") == "NOT_FOUND"
    assert normalize_yes_no_not_found("NOT_FOUND") == "NOT_FOUND"
    assert normalize_yes_no_not_found("not found") == "NOT_FOUND"
    assert normalize_yes_no_not_found("null") == "NOT_FOUND"
    assert normalize_yes_no_not_found("NOT_AVAILABLE") == "NOT_FOUND"
    assert normalize_yes_no_not_found("[blank]") == "NOT_FOUND"

    # Explicit positive
    assert normalize_yes_no_not_found("YES") == "YES"
    assert normalize_yes_no_not_found("yes") == "YES"
    assert normalize_yes_no_not_found("True") == "YES"

    # Explicit negative
    assert normalize_yes_no_not_found("NO") == "NO"
    assert normalize_yes_no_not_found("no") == "NO"
    assert normalize_yes_no_not_found("False") == "NO"
    assert normalize_yes_no_not_found("nil") == "NO"


# =============================================================================
# 4. EMIS ID Available Tests
# =============================================================================

def test_emis_id_positive_evidence_digits():
    doc_text = "Government Higher Secondary School\nEMIS ID: 2015743426\nTC No: 492"
    ans = evaluate_evidence_based_yes_no(
        field_key="is_emis_id_available",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
        page_number=1,
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.99
    assert ans["evidence"] == "EMIS ID: 2015743426"
    assert ans["source_document"] == "TRANSFER_CERTIFICATE"


def test_emis_id_positive_from_extracted_pool():
    # If emis_id was previously extracted in field pool
    ans = evaluate_evidence_based_yes_no(
        field_key="is_emis_id_available",
        document_text="",
        document_type="TRANSFER_CERTIFICATE",
        all_extracted_fields={"emis_id": "33020100101"},
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.99
    assert "33020100101" in ans["evidence"]


def test_emis_id_explicitly_nil_returns_no():
    doc_text = "Transfer Certificate\nEMIS ID: Nil\nAdmission No: 104"
    ans = evaluate_evidence_based_yes_no(
        field_key="is_emis_id_available",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "NO"
    assert ans["confidence"] >= 0.9
    assert "nil" in ans["evidence"].lower()


def test_emis_id_absent_without_explanation_returns_not_found():
    doc_text = "Transfer Certificate\nName: Kumar\nSchool: St Joseph High School"
    ans = evaluate_evidence_based_yes_no(
        field_key="is_emis_id_available",
        document_text=doc_text,
        document_type="TRANSFER_CERTIFICATE",
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["value"] != "NO"
    assert ans["confidence"] == 0.0
    assert ans["evidence"] is None


# =============================================================================
# 5. Derived Question Tests (Communication address same as permanent address)
# =============================================================================

def test_derived_address_same_identical():
    addr = "12/4 Gandhi Street, Anna Nagar, Chennai - 600040"
    pool = {
        "permanent_address": addr,
        "communication_address": addr,
    }
    ans = evaluate_evidence_based_yes_no(
        field_key="communication_address_same",
        all_extracted_fields=pool,
        document_type="AADHAAR",
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.95
    assert ans["evidence"] is not None


def test_derived_address_same_high_similarity():
    p_addr = "Plot 45, Second Street, Ram Nagar, Madurai 625001"
    c_addr = "45, 2nd Street, Ram Nagar, Madurai, Tamil Nadu 625001"
    pool = {
        "permanent_address": p_addr,
        "communication_address": c_addr,
    }
    ans = evaluate_evidence_based_yes_no(
        field_key="communication_address_same",
        all_extracted_fields=pool,
    )
    assert ans["value"] == "YES"
    assert ans["confidence"] >= 0.9


def test_derived_address_different_pincodes_returns_no():
    p_addr = "10 Nehru Street, Coimbatore 641001"
    c_addr = "25 Cross Road, Chennai 600028"
    pool = {
        "permanent_address": p_addr,
        "communication_address": c_addr,
    }
    ans = evaluate_evidence_based_yes_no(
        field_key="communication_address_same",
        all_extracted_fields=pool,
    )
    assert ans["value"] == "NO"
    assert ans["confidence"] >= 0.95
    assert "permanent" in ans["evidence"].lower() and "communication" in ans["evidence"].lower()


def test_derived_address_one_missing_returns_not_found():
    pool = {
        "permanent_address": "10 Nehru Street, Coimbatore 641001",
        # communication_address is missing
    }
    ans = evaluate_evidence_based_yes_no(
        field_key="communication_address_same",
        all_extracted_fields=pool,
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["value"] != "NO"  # Must NEVER infer NO because communication address is missing
    assert ans["confidence"] == 0.0
    assert ans["evidence"] is None


def test_derived_address_both_missing_returns_not_found():
    ans = evaluate_evidence_based_yes_no(
        field_key="communication_address_same",
        all_extracted_fields={},
    )
    assert ans["value"] == "NOT_FOUND"
    assert ans["confidence"] == 0.0


# =============================================================================
# 6. Accuracy Engine Integration & Structure Tests
# =============================================================================

def test_accuracy_engine_yes_no_structure():
    engine = AccuracyExtractionEngine()

    # Positive evidence test
    res_pos = engine.validate_field_accuracy(
        field_key="is_emis_id_available",
        raw_val=None,
        document_text="EMIS ID: 2015743426",
        source_doc_type="TRANSFER_CERTIFICATE",
        page_number=1,
    )
    assert res_pos.field_name == "is_emis_id_available"
    assert res_pos.normalized_value == "YES"
    assert res_pos.confidence == "HIGH"
    assert res_pos.confidence_score >= 0.9
    assert res_pos.field_state == "HIGH_CONFIDENCE"
    assert "2015743426" in res_pos.evidence_snippet

    d_pos = res_pos.to_dict()
    assert d_pos["is_available"] is True
    assert d_pos["answer"]["field"] == "is_emis_id_available"
    assert d_pos["answer"]["value"] == "YES"
    assert d_pos["answer"]["evidence"] == "EMIS ID: 2015743426"

    # Missing information test
    res_miss = engine.validate_field_accuracy(
        field_key="is_first_graduate",
        raw_val=None,
        document_text="Student Name: Anitha\nMarks: 450",
        source_doc_type="SSLC",
        page_number=1,
    )
    assert res_miss.field_name == "is_first_graduate"
    assert res_miss.normalized_value == "NOT_FOUND"
    assert res_miss.normalized_value != "NO"
    assert res_miss.confidence == "LOW"
    assert res_miss.confidence_score == 0.0
    assert res_miss.field_state == "NOT_FOUND"
    assert res_miss.evidence_snippet is None

    d_miss = res_miss.to_dict()
    assert d_miss["is_available"] is False
    assert d_miss["answer"]["value"] == "NOT_FOUND"
    assert d_miss["answer"]["confidence"] == 0.0
    assert d_miss["answer"]["evidence"] is None


# =============================================================================
# 7. Excel Writer & Column Alignment Tests (No Column Shifts)
# =============================================================================

def test_excel_writer_column_mapping_no_shift():
    """
    Ensures that Yes/No columns strictly receive YES, NO, or NOT_FOUND,
    and never shift values between columns.
    """
    service = FieldMappingService()
    excel_headers = [
        "Register No",
        "Student Name",
        "Is EMIS ID Available",
        "First Graduate",
        "Differently Abled",
        "Orphan",
        "Communication Address Same",
    ]

    mapped_fields = {
        "Register No": {"value": "9100201"},
        "Student Name": {"value": "Kavitha S"},
        "Is EMIS ID Available": {"value": "YES"},
        "First Graduate": {"value": "NO"},
        "Differently Abled": {"value": "NOT_FOUND"},
        "Orphan": {"value": None},  # unpopulated
        "Communication Address Same": {"value": "YES"},
    }

    write_dict = service.format_for_excel_write(excel_headers, mapped_fields)

    assert write_dict["Register No"] == "9100201"
    assert write_dict["Student Name"] in ("Kavitha S", "KAVITHA S")
    assert write_dict["Is EMIS ID Available"] == "YES"
    assert write_dict["First Graduate"] == "NO"
    assert write_dict["Differently Abled"] == "NOT_FOUND"
    assert write_dict["Orphan"] == "NOT_FOUND"  # Never NO when missing!
    assert write_dict["Communication Address Same"] == "YES"

    # Now verify openpyxl row write
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(excel_headers)

    row_data = [write_dict[h] for h in excel_headers]
    ws.append(row_data)

    # Check cell by cell
    assert ws.cell(row=2, column=1).value == "9100201"
    assert ws.cell(row=2, column=2).value in ("Kavitha S", "KAVITHA S")
    assert ws.cell(row=2, column=3).value == "YES"
    assert ws.cell(row=2, column=4).value == "NO"
    assert ws.cell(row=2, column=5).value == "NOT_FOUND"
    assert ws.cell(row=2, column=6).value == "NOT_FOUND"
    assert ws.cell(row=2, column=7).value == "YES"


def test_format_for_excel_write_standalone():
    assert format_for_excel_write("Is First Graduate?", "yes") == "YES"
    assert format_for_excel_write("Is First Graduate?", "no") == "NO"
    assert format_for_excel_write("Is First Graduate?", None) == "NOT_FOUND"
    assert format_for_excel_write("Is First Graduate?", "") == "NOT_FOUND"
    assert format_for_excel_write("Is First Graduate?", "not_found") == "NOT_FOUND"
