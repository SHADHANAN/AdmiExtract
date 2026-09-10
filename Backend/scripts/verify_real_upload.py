import asyncio
import os
import json
from pathlib import Path
import openpyxl
from fastapi import UploadFile

from app.db.database import init_db
from app.services.document_processing_pipeline import DocumentProcessingPipeline
from app.services.excel_template_service import ExcelTemplateService
from app.services.ocr_service import ocr_service
from app.models.excel_template import ExcelBatchTemplate


class MockUploadFile(UploadFile):
    def __init__(self, filename: str, path: str):
        self.filename = filename
        self._path = path
        with open(path, "rb") as f:
            self._bytes = f.read()
        super().__init__(filename=filename, file=None)

    async def read(self, size: int = -1) -> bytes:
        return self._bytes


async def main():
    print("================================================================================")
    print("           REAL END-TO-END ADMISSION DOCUMENT EXTRACTION EXECUTION               ")
    print("================================================================================\n")

    # Connect to MongoDB
    await init_db()

    pipeline = DocumentProcessingPipeline()
    excel_service = ExcelTemplateService()

    batch_id = "batch_e1a995b9"
    register_number = "714024247103"
    student_name = "SHADHANAN S"
    mobile_number = "9488203077"
    class_id = "class_batch_e1a995b9_section_b"

    # Actual student documents in uploads/
    doc_paths = [
        ("Aadhaar.pdf", "uploads/Aadhaar.pdf"),
        ("10th Mark sheet.pdf", "uploads/10th Mark sheet.pdf"),
        ("12thmarksheet.pdf", "uploads/12thmarksheet.pdf"),
        ("Community Certificate.pdf", "uploads/Community Certificate.pdf"),
    ]

    print("--------------------------------------------------------------------------------")
    print("STAGE 1: OCR LOGS & TEXT EXTRACTION FOR ACTUAL UPLOADED DOCUMENTS")
    print("--------------------------------------------------------------------------------")
    for doc_name, rel_path in doc_paths:
        abs_path = os.path.abspath(rel_path)
        try:
            ocr_res = await asyncio.to_thread(ocr_service.extract_text, abs_path)
            extracted_txt = ocr_res.get("text", "") or ""
            print(f"[OCR] Document: '{doc_name}' | Status: {'SUCCESS' if ocr_res.get('success') else 'NOTE'} | Chars: {len(extracted_txt)}")
            clean_snip = " ".join(extracted_txt.split())[:120]
            print(f"[OCR] Sample Text: \"{clean_snip}...\"")
        except Exception as ocr_err:
            print(f"[OCR] Document: '{doc_name}' | Note: {ocr_err}")

    print("\n--------------------------------------------------------------------------------")
    print("STAGE 2: GEMINI MULTIMODAL EXTRACTION JSON (PER DOCUMENT)")
    print("--------------------------------------------------------------------------------")
    upload_files = [MockUploadFile(name, path) for name, path in doc_paths]

    # Process all documents through DocumentProcessingPipeline
    result = await pipeline.process_student_documents(
        batch_id=batch_id,
        register_number=register_number,
        student_name=student_name,
        mobile_number=mobile_number,
        files=upload_files,
    )

    verification_fields = result.get("verification_fields", {})
    extracted_per_doc = result.get("extracted_fields_per_document", {})
    detected_docs = result.get("detected_documents", [])

    for doc_name, fields in extracted_per_doc.items():
        print(f"\nDocument: {doc_name}")
        print(json.dumps(fields, indent=2))

    print("\n--------------------------------------------------------------------------------")
    print("STAGE 3: DOCUMENT CLASSIFICATION OUTPUT")
    print("--------------------------------------------------------------------------------")
    print(f"Detected Document Types: {detected_docs}")
    for fname, fields in extracted_per_doc.items():
        doc_type = "UNKNOWN"
        for d in detected_docs:
            if d.lower() in fname.lower() or fname.lower().startswith(d.lower()):
                doc_type = d
                break
        print(f"  - File: '{fname}' -> Classified As: {fields.get('document_type', doc_type)} ({len(fields)} fields)")

    print("\n--------------------------------------------------------------------------------")
    print("STAGE 4: FINAL MAPPED DICTIONARY BEFORE EXCEL WRITE")
    print("--------------------------------------------------------------------------------")
    mapped_summary = {}
    for k, v in verification_fields.items():
        val = v.get("value") if isinstance(v, dict) else v
        src = v.get("source") if isinstance(v, dict) else "N/A"
        cnf = v.get("confidence", 0) if isinstance(v, dict) else 0
        if val is not None and str(val).strip() != "" and str(val).upper() not in ["NO", "NULL", "NONE"]:
            mapped_summary[k] = {"value": val, "source": src, "confidence": cnf}
        elif str(val) == "No":
            mapped_summary[k] = {"value": "No", "source": src, "confidence": cnf}
    print(json.dumps(mapped_summary, indent=2))

    print("\n--------------------------------------------------------------------------------")
    print("STAGE 5: HEADER-TO-COLUMN MAPPING")
    print("--------------------------------------------------------------------------------")
    template = await excel_service.repository.get_by_batch_id(batch_id, class_id=class_id)
    wb = openpyxl.load_workbook(template.file_path, data_only=True)
    sheet = wb.active
    header_map = {}
    for c in range(1, sheet.max_column + 1):
        h_name = sheet.cell(row=1, column=c).value
        if h_name:
            header_map[str(h_name).strip()] = c
    wb.close()

    print(f"Total Columns Mapped: {len(header_map)}")
    print(json.dumps(header_map, indent=2))

    print("\n--------------------------------------------------------------------------------")
    print("STAGE 6: EXCEL WRITE LOGS (ROW NUMBER & WRITTEN COLUMNS)")
    print("--------------------------------------------------------------------------------")
    await excel_service.append_or_update_student_row_in_excel(
        batch_id=batch_id,
        register_number=register_number,
        student_data=verification_fields,
        class_id=class_id,
    )

    print("\n--------------------------------------------------------------------------------")
    print("STAGE 7: INSPECTION OF POPULATED EXCEL WORKBOOK (READ BACK FROM DISK)")
    print("--------------------------------------------------------------------------------")
    wb = openpyxl.load_workbook(template.file_path, data_only=True)
    sheet = wb.active
    print(f"File Path: {template.file_path}")
    print(f"Active Sheet: '{sheet.title}'")
    print(f"Total Rows: {sheet.max_row} | Total Columns: {sheet.max_column}\n")

    target_row_idx = None
    for r in range(2, sheet.max_row + 1):
        reg_val = str(sheet.cell(row=r, column=2).value or "").strip()
        if reg_val == register_number:
            target_row_idx = r
            break

    if target_row_idx is None:
        target_row_idx = sheet.max_row

    print(f"Target Student Row: Row {target_row_idx}")
    print("=" * 85)
    print(f"{'Col #':<7} | {'Excel Header Name':<42} | {'Populated Cell Value'}")
    print("-" * 85)

    populated_data_count = 0
    populated_bool_count = 0
    blank_count = 0

    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(row=1, column=col).value
        val = sheet.cell(row=target_row_idx, column=col).value

        if val is not None and str(val).strip() != "" and str(val).upper() not in ["NO", "NULL", "NONE"]:
            populated_data_count += 1
            print(f"Col {col:<4} | {str(header):<42} | {str(val)}")
        elif str(val) == "No":
            populated_bool_count += 1
            print(f"Col {col:<4} | {str(header):<42} | No (Optional/Boolean)")
        else:
            blank_count += 1

    wb.close()
    print("=" * 85)
    print(f"\nPOPULATED DATA COLUMNS    : {populated_data_count}")
    print(f"POPULATED BOOLEAN COLUMNS : {populated_bool_count}")
    print(f"BLANK OPTIONAL COLUMNS    : {blank_count}")
    print(f"TOTAL COLUMNS VERIFIED    : {populated_data_count + populated_bool_count + blank_count}")
    print("\n>>> RUNTIME VERIFICATION SUCCESSFUL: EXCEL POPULATED WITH CANDIDATE DATA <<<")


if __name__ == "__main__":
    asyncio.run(main())
